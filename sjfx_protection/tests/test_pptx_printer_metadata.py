"""Real python-pptx print metadata recognition remains narrow and fail closed."""
import hashlib
import io
import struct
import unittest
import warnings
import xml.etree.ElementTree as ET
import zipfile

from pptx import Presentation

from sjfx_local_protection.risk import inspect_bytes


PART = "ppt/printerSettings/printerSettings1.bin"
TYPES = "[Content_Types].xml"
RELS = "ppt/_rels/presentation.xml.rels"
CT_NS = "{http://schemas.openxmlformats.org/package/2006/content-types}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.printerSettings"
REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/printerSettings"


class PrinterMetadataTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Synthetic ordinary presentation"
        slide.placeholders[1].text = "Budget 125000 yuan. 项目预算为125000元。"
        output = io.BytesIO()
        presentation.save(output)
        cls.original = output.getvalue()
        with zipfile.ZipFile(io.BytesIO(cls.original)) as archive:
            cls.entries = [(item.filename, archive.read(item)) for item in archive.infolist()]
        cls.profile = dict(cls.entries)[PART]
        assert len(cls.profile) == 9395
        assert hashlib.sha256(cls.profile).hexdigest() == "d7768f87e07d29634782e448ece1cddd05a52b9499254222b1190bc8f6dc579e"

    def package(self, *, replacements=None, deleted=(), additions=(), rename=None, printer_first=False):
        replacements = replacements or {}
        rename = rename or {}
        entries = [(rename.get(name, name), replacements.get(name, data))
                   for name, data in self.entries if name not in deleted]
        entries += list(additions)
        if printer_first:
            entries.sort(key=lambda item: 0 if item[0] == PART else 1)
        output = io.BytesIO()
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                for name, data in entries:
                    # ZipInfo normally normalizes Windows backslashes at
                    # construction. Write the requested raw archive name so
                    # noncanonical-name tests exercise actual ZIP bytes.
                    item = zipfile.ZipInfo(name)
                    item.filename = item.orig_filename = name
                    archive.writestr(item, data, compress_type=zipfile.ZIP_DEFLATED)
        return output.getvalue()

    def inspect(self, raw=None, name="presentation.pptx", **policy):
        result = inspect_bytes(self.original if raw is None else raw, name,
                               {"required_av": False, **policy})
        self.assertEqual(result["engines"]["av"]["status"], "not_requested")
        return result

    def assert_incomplete(self, raw, *, reason="office_binary_part_requires_parser", **policy):
        result = self.inspect(raw, **policy)
        self.assertFalse(result["complete"], result)
        self.assertIn(reason, result["engines"]["static"]["incomplete_reasons"])
        self.assertNotEqual(result["verdict"], "no_detection")
        return result

    def xml_change(self, name, action):
        root = ET.fromstring(dict(self.entries)[name])
        action(root)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    def test_real_generated_pptx_complete_and_every_profile_byte_evidenced(self):
        result = self.inspect()
        self.assertEqual(result["verdict"], "no_detection", result)
        self.assertTrue(result["complete"])
        self.assertIn("office_known_printer_metadata", result["engines"]["static"]["formats"])
        evidence = next(item for item in result["members"] if item["member_name"] == PART)
        self.assertEqual(evidence["size"], len(self.profile))
        self.assertEqual(evidence["sha256"], hashlib.sha256(self.profile).hexdigest())
        # Reading through python-pptx preserves the ordinary slide text; the
        # profile remains a raw part instead of invoking a driver/parser.
        reopened = Presentation(io.BytesIO(self.original))
        self.assertIn("125000", reopened.slides[0].placeholders[1].text)
        printer = reopened.part.part_related_by(REL_TYPE)
        self.assertEqual(printer.blob, self.profile)

    def test_zip_member_order_and_valid_content_type_override(self):
        def change(types):
            for elem in types:
                if elem.attrib.get("Extension") == "bin":
                    elem.set("ContentType", "application/octet-stream")
            ET.SubElement(types, CT_NS + "Override", PartName="/" + PART, ContentType=CONTENT_TYPE)
        raw = self.package(replacements={TYPES: self.xml_change(TYPES, change)}, printer_first=True)
        self.assertEqual(self.inspect(raw)["verdict"], "no_detection")

    def test_changed_truncated_appended_and_other_driver_profiles_incomplete(self):
        mutated = self.profile[:-1] + bytes([self.profile[-1] ^ 1])
        devmode = bytearray(220)
        struct.pack_into("<HHHHI", devmode, 64, 0x0401, 1, 220, 0, 0)
        for profile in (mutated, self.profile[:-1], self.profile + b"extra", bytes(devmode), b"unknown binary"):
            with self.subTest(size=len(profile)):
                self.assert_incomplete(self.package(replacements={PART: profile}))

    def test_content_type_is_required_and_cannot_be_ambiguous(self):
        def wrong(types):
            for elem in types:
                if elem.attrib.get("Extension") == "bin":
                    elem.set("ContentType", "application/octet-stream")
        def duplicate(types):
            ET.SubElement(types, CT_NS + "Default", Extension="bin", ContentType=CONTENT_TYPE)
        def override_wrong(types):
            ET.SubElement(types, CT_NS + "Override", PartName="/" + PART, ContentType="application/octet-stream")
        for change in (wrong, duplicate, override_wrong):
            with self.subTest(change=change.__name__):
                self.assert_incomplete(self.package(replacements={TYPES: self.xml_change(TYPES, change)}))
        self.assert_incomplete(self.package(deleted=(TYPES,)), reason="unsupported_format")

    def test_relationship_is_required_internal_canonical_and_unambiguous(self):
        for key, value in (("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject"),
                           ("TargetMode", "External"), ("Target", "./printerSettings/printerSettings1.bin"),
                           ("Target", "/ppt/printerSettings/printerSettings1.bin")):
            def change(root, key=key, value=value):
                next(e for e in root if e.attrib.get("Type") == REL_TYPE).set(key, value)
            with self.subTest(attribute=key, value=value):
                self.assert_incomplete(self.package(replacements={RELS: self.xml_change(RELS, change)}))
        def remove(root):
            root.remove(next(e for e in root if e.attrib.get("Type") == REL_TYPE))
        def duplicate(root):
            existing = next(e for e in root if e.attrib.get("Type") == REL_TYPE)
            ET.SubElement(root, REL_NS + "Relationship", **existing.attrib)
        for change in (remove, duplicate):
            self.assert_incomplete(self.package(replacements={RELS: self.xml_change(RELS, change)}))
        self.assert_incomplete(self.package(deleted=(RELS,)))

    def test_canonical_member_and_pptx_container_required_even_for_known_bytes(self):
        for target in ("ppt/printerSettings/other.bin", "ppt\\printerSettings\\printerSettings1.bin",
                       "ppt/embeddings/printerSettings1.bin"):
            with self.subTest(target=target):
                self.assert_incomplete(self.package(rename={PART: target}))
        standalone = self.inspect(self.profile, name=PART)
        self.assertFalse(standalone["complete"])
        other_office = self.package(additions=(("word/document.xml", b"<document/>"),))
        result = self.inspect(other_office, name="document.docx")
        self.assertIn("office_binary_part_requires_parser", result["engines"]["static"]["incomplete_reasons"])

    def test_unknown_vba_ole_and_executable_binary_not_admitted(self):
        for name, data in (("ppt/unknown.bin", self.profile), ("ppt/vbaProject.bin", b"opaque macro bytes"),
                           ("ppt/embeddings/object.bin", bytes.fromhex("d0cf11e0a1b11ae1") + b"opaque OLE")):
            with self.subTest(name=name):
                result = self.inspect(self.package(additions=((name, data),)))
                self.assertFalse(result["complete"])
                self.assertNotEqual(result["verdict"], "no_detection")
        for magic in (b"MZ", b"\x7fELF"):
            result = self.inspect(self.package(replacements={PART: magic + self.profile[len(magic):]}))
            self.assertEqual(result["verdict"], "dangerous")
            self.assertIn("disguised_executable", {f["rule"] for f in result["findings"]})

    def test_known_metadata_never_bypasses_budget_or_duplicate_part_checks(self):
        self.assert_incomplete(self.original, reason="member_limit", max_members=1)
        self.assert_incomplete(self.original, reason="expansion_limit", max_expanded_bytes=9394)
        self.assert_incomplete(self.package(additions=((PART, self.profile),)), reason="duplicate_archive_member")

    def test_metadata_parser_is_bounded_and_entities_fail_closed(self):
        padded = dict(self.entries)[TYPES] + b" " * (64 * 1024)
        self.assert_incomplete(self.package(replacements={TYPES: padded}))
        declared = '<?xml version="1.0"?><!DOCTYPE Types [<!ENTITY x "unused">]>'
        original = dict(self.entries)[TYPES].decode("utf-8")
        original = original[original.index("<Types"):]
        for encoding in ("utf-8", "utf-16"):
            result = self.assert_incomplete(self.package(replacements={TYPES: (declared + original).encode(encoding)}))
            self.assertIn("office_xml_entity", result["engines"]["static"]["incomplete_reasons"])


if __name__ == "__main__":
    unittest.main()
