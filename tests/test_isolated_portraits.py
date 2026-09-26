"""Behavioral regression tests for isolated, bounded content recognition."""
import json
import os
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from config import Config
from services.large_package_isolated import LargePackageStore, iter_inventory
from services.large_package_preview import preview_file, portrait
from services.large_package_preview_runner import PreviewRunner, PreviewCancelled
from large_package_worker import LargePackageWorker


class IsolatedPortraitTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.allowed = patch.object(Config, 'SCAN_ALLOWED_ROOTS', ())
        self.allowed.start()
        self.catalog_step = patch.object(LargePackageWorker, "_catalog_incremental")
        self.catalog_step.start()

    def tearDown(self):
        self.catalog_step.stop()
        self.allowed.stop()
        self.tmp.cleanup()

    def item(self, name, text):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding='utf-8')
        stat = path.stat()
        return {'path': name, 'name': path.name, 'extension': path.suffix, 'size': stat.st_size, 'modified_at_ns': stat.st_mtime_ns}

    def preview(self, name, text, refine=False):
        row = self.item(name, text)
        return preview_file(self.root, row, refine)[0]

    def store(self):
        store = LargePackageStore(self.root / 'state' / 'large.db')
        package = store.create(self.root)
        return store, package

    def test_same_format_different_business(self):
        sales = self.preview('a.csv', '订单日期,客户,销售额,回款状态\n2026-09-01,甲公司,100,未回款\n')
        repair = self.preview('b.csv', '设备编号,故障类型,检修时间,维修人员\nA01,报警,2026-09-01,张三\n')
        self.assertEqual(sales['portrait']['topic'], '销售与回款')
        self.assertEqual(repair['portrait']['topic'], '设备故障与维护')
        self.assertTrue(sales['portrait']['evidence'])
        self.assertFalse(sales['complete'])
        self.assertEqual(sales['source_hash_kind'], 'pending')

    def test_unfamiliar_schema_keeps_specific_fields(self):
        value = self.preview('unknown.csv', '星体编号,光谱频率,观测台站\nA1,125,北站\n')
        self.assertIn('光谱频率', value['portrait']['topic'])
        self.assertTrue(value['portrait']['needs_refine'])

    def test_names_cannot_determine_content_topic(self):
        value = self.preview('销售回款.csv', '设备编号,故障类型,检修时间\nA1,报警,2026\n')
        self.assertEqual(value['portrait']['topic'], '设备故障与维护')

    def test_ambiguous_topics_remain_reviewable(self):
        info = portrait({'path': 'x.txt'}, [{'location':'段落', 'text':'销售额 回款 故障 维修', 'kind':'text'}], [], [])
        self.assertTrue(info['ambiguous'])
        self.assertTrue(info['needs_refine'])
        self.assertNotEqual(info['quality'], 'identified')

    def test_sparse_large_file_does_not_hash_full_source(self):
        path = self.root / 'big.txt'
        with path.open('wb') as f:
            f.write(b'Head sales receivable\n')
            f.seek(2 * 1024 ** 3)
            f.write(b'\nTail maintenance fault\n')
        stat = path.stat()
        row = {'path':path.name, 'size':stat.st_size, 'modified_at_ns':stat.st_mtime_ns}
        before = time.monotonic()
        summary, source_hash = preview_file(self.root, row)
        self.assertIsNone(source_hash)
        self.assertLessEqual(summary['sampled_bytes'], 65536)
        self.assertLess(time.monotonic() - before, 3)

    def test_changed_source_rejected(self):
        row = self.item('change.txt', 'before')
        (self.root / 'change.txt').write_text('changed size', encoding='utf-8')
        with self.assertRaises(RuntimeError):
            preview_file(self.root, row)

    def test_root_escape_rejected(self):
        with self.assertRaises(ValueError):
            preview_file(self.root, {'path':'../escape.txt'})

    def test_sensitive_not_read(self):
        row = self.item('ordinary.txt', 'private content')
        row['sensitive'] = True
        value, digest = preview_file(self.root, row)
        self.assertEqual(value['status'], 'restricted')
        self.assertEqual(value['sample'], '')
        self.assertIsNone(digest)

    def test_json_records_and_fields(self):
        value = self.preview('rows.json', json.dumps([{'销售额':100, '回款状态':'未回款'}], ensure_ascii=False))
        # JSON recognition preserves field types, unlike plain CSV headers.
        self.assertEqual(value['portrait']['fields'], ['销售额:int', '回款状态:str'])
        self.assertEqual(value['portrait']['topic'], '销售与回款')

    def test_email_headers_are_decoded(self):
        value = self.preview('message.eml', 'Subject: Sales receivable\nFrom: a@example.com\nTo: b@example.com\nContent-Type: text/plain; charset=utf-8\n\nSales receivable remains open.')
        self.assertEqual(value['portrait']['document_role'], '邮件')
        self.assertIn('a@example.com', value['portrait']['entities'])

    def test_docx_content_not_style_xml(self):
        from docx import Document
        path = self.root / 'memo.docx'
        doc = Document()
        doc.add_heading('设备故障与检修', 0)
        doc.add_paragraph('设备编号 A01。故障报警后安排维修。')
        doc.save(path)
        value, _ = preview_file(self.root, {'path':path.name})
        self.assertEqual(value['portrait']['topic'], '设备故障与维护')
        self.assertNotIn('w:style', value['sample'])

    def test_excel_sheet_fields_distinguish_content(self):
        from openpyxl import Workbook
        path = self.root / 'book.xlsx'
        book = Workbook()
        book.active.title = '销售明细'
        book.active.append(['订单日期','客户','销售额','回款状态'])
        book.active.append(['2026-09-01','甲',100,'未回款'])
        book.save(path)
        book.close()
        value, _ = preview_file(self.root, {'path':path.name})
        self.assertIn('销售明细', value['portrait']['headings'])
        self.assertIn('销售额', value['portrait']['fields'])
        self.assertEqual(value['portrait']['topic'], '销售与回款')

    def test_ppt_samples_middle_slide(self):
        from pptx import Presentation
        path = self.root / 'deck.pptx'
        pres = Presentation()
        for text in ('封面','项目计划 交付 里程碑','结束'):
            slide = pres.slides.add_slide(pres.slide_layouts[1])
            slide.shapes.title.text = text
        pres.save(path)
        value, _ = preview_file(self.root, {'path':path.name})
        self.assertIn('里程碑', value['sample'])
        self.assertEqual(value['portrait']['topic'], '项目进度与交付')

    def test_empty_pdf_deferred_for_ocr(self):
        from pypdf import PdfWriter
        path = self.root / 'blank.pdf'
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(path)
        value, _ = preview_file(self.root, {'path':path.name})
        self.assertEqual(value['portrait']['recognition_status'], 'needs_ocr')
        self.assertFalse(value['coverage']['complete'])

    def test_queue_selection_preserves_previews_and_deep_results(self):
        store, package = self.store()
        row = self.item('data.csv', '销售额,回款\n100,未回款\n')
        store.upsert_file(package['id'], row)
        value, _ = preview_file(self.root, row)
        worker = LargePackageWorker(store)
        worker._save_preview(package['id'], row['path'], value)
        store.update_file(package['id'], row['path'], deep_status='completed', deep_summary={'summary_type':'model','summary':'existing'})
        store.update_package(package['id'],status='waiting_for_selection',phase='catalog_ready')
        store.save_selection(package['id'], [row['path']])
        current = store.get_file(package['id'], row['path'])
        self.assertEqual(current['deep_summary']['summary'], 'existing')
        self.assertEqual(current['deep_status'], 'queued')
        self.assertIn('销售额', current['local_summary']['portrait']['topic'])
        self.assertEqual(current['local_summary']['portrait']['catalog_status'], 'pending')

    def test_failed_refinement_preserves_first_pass(self):
        store, package = self.store()
        row = self.item('data.csv', '星体编号,光谱频率\nA1,100\n')
        store.upsert_file(package['id'], row)
        worker = LargePackageWorker(store)
        worker._save_preview(package['id'], row['path'], preview_file(self.root, row)[0])
        worker._save_preview(package['id'], row['path'], {'sample':'','portrait':{},'notice':'timeout'}, refine=True)
        current = store.get_file(package['id'], row['path'])
        self.assertIn('光谱频率', current['local_summary']['sample'])
        self.assertEqual(current['recognition_status'], 'needs_review')

    def test_pages_reach_all_files_and_remaining_is_server_filtered(self):
        store, package = self.store()
        store.upsert_files(package['id'], [{'path':str(i)+'.csv','name':str(i)+'.csv','size':1} for i in range(305)])
        page = store.file_page(package['id'], limit=100, offset=300)
        self.assertEqual(page['total'], 305)
        self.assertEqual(len(page['files']), 5)
        for row in store.file_page(package['id'], limit=300)['files']:
            store.update_file(package['id'],row['path'],deep_status='completed')
        self.assertEqual(store.file_page(package['id'],remaining=True)['total'], 5)

    def test_catalog_has_topics_and_facets_without_touching_ordinary_db(self):
        store, package = self.store()
        ordinary = self.root / 'agent.db'
        ordinary.write_bytes(b'ordinary sentinel')
        worker = LargePackageWorker(store)
        for name, body in [('a.csv','销售额,回款状态\n100,待回款\n'),('b.csv','设备编号,故障,维修\nA01,报警,完成\n')]:
            row=self.item(name,body);store.upsert_file(package['id'],row)
            worker._save_preview(package['id'],name,preview_file(self.root,row)[0])
        catalog=store.catalog(package['id'])
        self.assertEqual(len(catalog['smart']['topics']),2)
        self.assertEqual(catalog['smart']['catalog_states']['pending'],2)
        self.assertEqual(ordinary.read_bytes(),b'ordinary sentinel')
        topic=catalog['smart']['topics'][0]['topic']
        self.assertEqual(store.file_page(package['id'],topic=topic)['total'],1)
        self.assertEqual(store.file_page(package['id'],facet_kind='role',facet_value='数据表')['total'],2)

    def test_legacy_migration_and_explicit_upgrade_preserve_results(self):
        store, package = self.store()
        store.upsert_file(package['id'],self.item('legacy.txt','content'))
        store.update_file(package['id'],'legacy.txt',deep_status='completed',deep_summary={'summary':'keep'})
        store.update_package(package['id'],status='completed',phase='report_ready')
        self.assertEqual(store.request_recognition(package['id']),1)
        self.assertEqual(store.get(package['id'])['phase'],'recognize_completed')
        self.assertEqual(store.get_file(package['id'],'legacy.txt')['deep_summary']['summary'],'keep')

    def test_no_requeue_of_finished_refinement_while_pending_remain(self):
        store, package=self.store()
        for name in ['one.txt','two.txt']:
            store.upsert_file(package['id'],self.item(name,'some content'))
            store.enqueue_recognition(package['id'],name)
        store.finish_queue(package['id'],'one.txt','recognize',True)
        store.update_file(package['id'],'one.txt',recognition_status='needs_review')
        store.update_package(package['id'],status='waiting_for_selection',phase='catalog_ready')
        self.assertEqual(store.request_recognition(package['id']),1)
        self.assertEqual(store.stage_counts(package['id'],'recognize').get('completed'),1)

    def test_real_runner_cancellation(self):
        runner=PreviewRunner()
        try:
            with self.assertRaises(PreviewCancelled):
                runner.run(self.root,self.item('cancel.txt','test'),cancel_check=lambda:True)
            self.assertIsNone(runner.process)
        finally:
            runner.close()

    def test_real_worker_lifecycle(self):
        data=self.root/'input';data.mkdir()
        self.item('input/a.csv','销售额,回款状态\n100,待回款\n')
        store=LargePackageStore(self.root/'state'/'large.db');package=store.create(data)
        worker=LargePackageWorker(store)
        try:
            worker.run_once()
            original_counts = store.stage_counts
            with patch.object(store, 'stage_counts', side_effect=lambda pid, stage: {} if stage == 'catalog' else original_counts(pid, stage)):
                worker.run_once();worker.run_once()
            self.assertEqual(store.get(package['id'])['status'],'waiting_for_selection')
            self.assertEqual(store.counts(package['id'])['recognition']['needs_review'],1)
            store.update_package(package['id'],status='paused')
            self.assertFalse(worker.run_once())
        finally:
            worker.preview_runner.close()

    def test_refinement_global_budget_keeps_remaining_queue(self):
        store, package = self.store()
        for name in ('one.txt', 'two.txt', 'three.txt'):
            row=self.item(name, '星体观测 频率记录')
            store.upsert_file(package['id'], row)
            store.enqueue_recognition(package['id'], name)
        store.update_package(package['id'], status='running', phase='recognize')
        worker=LargePackageWorker(store)
        try:
            with patch.dict(os.environ, {'SJFX_LP_RECOGNITION_SLICE_FILES':'1'}):
                worker._recognize_one(store.get(package['id']))
            self.assertEqual(store.get(package['id'])['status'], 'waiting_for_selection')
            self.assertEqual(store.stage_counts(package['id'],'recognize')['queued'],2)
        finally:
            worker.preview_runner.close()

    def test_old_database_migration_is_additive(self):
        import sqlite3
        db=self.root/'old.db'
        conn=sqlite3.connect(db)
        conn.execute("CREATE TABLE files(package_id TEXT,path TEXT,name TEXT,extension TEXT,size INTEGER,modified_at_ns INTEGER,sha256 TEXT,sensitive INTEGER,file_type TEXT,quick_status TEXT,deep_status TEXT,category TEXT,confidence REAL,local_summary TEXT,deep_summary TEXT,error TEXT,updated_at REAL,PRIMARY KEY(package_id,path))")
        conn.execute("INSERT INTO files VALUES('lp-old','a.txt','a.txt','.txt',1,1,'oldhash',0,'text','completed','completed','文本资料',0.5,'{}','{\"summary\":\"keep\"}',NULL,1)")
        conn.commit();conn.close()
        store=LargePackageStore(db)
        row=store.get_file('lp-old','a.txt')
        self.assertEqual(row['sha256'],'oldhash')
        self.assertEqual(row['deep_summary']['summary'],'keep')
        self.assertEqual(row['recognition_status'],'legacy')

    def test_deep_parser_metadata_only_is_not_completed(self):
        worker=LargePackageWorker()
        parser=unittest.mock.Mock()
        parser.parse.return_value={'text':'','evidence':[],'warnings':['over size']}
        worker._parser=parser
        result=worker._deep_parse(self.root, {'path':'big.pdf'})
        self.assertEqual(result['summary_type'],'local_deep_failed')

    def test_api_paging_filters_and_owner_boundary(self):
        import app as application
        store, package=self.store()
        store.upsert_files(package['id'],[{'path':str(i)+'.txt','name':str(i)+'.txt'} for i in range(205)])
        with patch.object(application,'large_package_store',store), patch.object(application,'_request_owner_id',return_value='legacy'):
            with application.app.test_request_context('/api/large-packages/x/files?offset=200&limit=100'):
                payload=application.list_large_package_files(package['id']).get_json()
                self.assertEqual(payload['total'],205)
                self.assertEqual(len(payload['files']),5)
            store.update_package(package['id'],status='waiting_for_selection',phase='catalog_ready')
            with application.app.test_request_context('/api/large-packages/x/recognize',method='POST',json={}):
                response, status=application.recognize_large_package(package['id'])
                self.assertEqual(status,202)
                self.assertEqual(response.get_json()['queued_files'],205)
        with patch.object(application,'large_package_store',store), patch.object(application,'_request_owner_id',return_value='someone-else'):
            with application.app.test_request_context('/api/large-packages/x/files'):
                response,status=application.list_large_package_files(package['id'])
                self.assertEqual(status,404)

    def test_report_and_exports_keep_coverage(self):
        from services.large_package_reporting import build_report, export_report
        store,package=self.store();row=self.item('report.csv','销售额,回款状态\n100,未回款\n')
        store.upsert_file(package['id'],row)
        worker=LargePackageWorker(store)
        worker._save_preview(package['id'],row['path'],preview_file(self.root,row)[0])
        store.update_file(package['id'],row['path'],deep_status='completed',deep_summary={'summary_type':'model','summary':'报告测试','text_preview':'销售额与回款待核验'})
        store.update_package(package['id'],status='completed',phase='report_ready')
        item=store.get(package['id'])
        report=build_report(store,item)
        self.assertEqual(report['recognition']['needs_review'],1)
        self.assertEqual(report['coverage']['deep_completed'],1)
        self.assertTrue(export_report(store,item,'json').is_file())
        self.assertTrue(export_report(store,item,'docx').is_file())



if __name__ == '__main__':
    unittest.main()
