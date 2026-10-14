"""Temporary DOCX/XML contract checks, without rendering or application setup."""
import ast
import copy
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from docx import Document

from services.exporter import create_report_docx


SCAN = {"root": "/fixture", "file_count": 2, "directory_count": 0,
        "total_size_human": "2 KB", "scanned_at": "2026-10-11"}


def report_fixture():
    counter = {"evidence_id": "C-last", "source_path": "papers/failure-study.pdf", "page": 7,
               "section": "失败条件", "doi": "10.0000/failure-fixture", "support_status": "supported",
               "text": "原文逐项比较了模型条件与原基线。" * 24 + "不能改善准确率|不适用于噪声标准差>0.4的样本。"}
    support = {"evidence_id": "S-method", "source_path": "papers/method-study.pdf", "page": 3,
               "section": "方法", "text": "已有方法在清晰输入上比较了准确率，尚未验证噪声条件。"}
    # The limiting record is deliberately after the old 12-card cutoff.
    background = [{"evidence_id": "B-%s" % i, "source_path": "papers/background.pdf", "text": "资料中的背景条件%s。" % i} for i in range(12)]
    return {
        "research_version": "ordinary-opportunities/1.0", "generation_mode": "model_analyzed",
        "basic_information": ["已选择两份资料"], "key_findings": ["清晰条件已有结果，噪声条件待复核。"],
        "recommended_research_direction": {
            "title": "验证噪声条件下的适用边界", "priority": "中", "confidence": "低",
            "rationale": "资料揭示了适用条件与失败结果，先开展边界核验。",
            "selection_reason": "优先核验已有失败条件，强备选需要额外数据。",
            "known_work": "方法论文已经比较清晰输入及原基线。",
            "research_gap": "噪声输入的资源条件和评价口径仍待核验。",
            "hypothesis": "噪声条件下保持可靠性是否可行？", "hypothesis_status": "unverified",
            "research_questions": ["噪声条件下保持可靠性是否可行？"],
            "methods": ["核对原文方法与失败条件 → 条件对照表"],
            "quality_gate": {"status": "passed", "resources": "confirmation_first"},
            "evidence_chain": [support] + background + [counter], "counterevidence": [counter],
            "evidence_roles": {"support": ["S-method", "C-last"], "prior_work": ["S-method"],
                               "counterevidence": ["C-last"], "conditions": ["C-last"]},
            "limitations": ["仅比较当前导入资料；来源身份不同不代表独立实验。"],
        },
    }


def generated(tmp_path, report, scan=None, name="fixture.docx"):
    output = tmp_path / name
    create_report_docx(report, scan or SCAN, output)
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml")
        root = ElementTree.fromstring(xml)
        namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs = ["".join(p.itertext()) for p in root.findall(".//w:p", namespaces)]
        styles = archive.read("word/styles.xml")
    return output, "\n".join(paragraphs), xml, styles


def test_new_docx_keeps_page_research_fields_and_full_late_counterevidence(tmp_path):
    report = report_fixture()
    before = copy.deepcopy(report)
    output, text, _, _ = generated(tmp_path, report)
    recommendation = report["recommended_research_direction"]
    for field in ("selection_reason", "known_work", "research_gap", "hypothesis"):
        assert recommendation[field] in text
    assert "假设状态：尚未验证。" in text
    assert "先核验资源" in text
    assert recommendation["counterevidence"][0]["text"] in text
    assert "[C-last] 来源：papers/failure-study.pdf，第 7 页，失败条件；角色：反证线索、适用条件" in text
    assert "来源 DOI：10.0000/failure-fixture" in text
    assert "[S-method] 来源：papers/method-study.pdf，第 3 页，方法；角色：已有工作、背景支撑" in text
    counter = text.split("[C-last]", 1)[1].split("附件二", 1)[0]
    assert "背景支撑" not in counter
    assert text.count("[C-last]") == 1
    assert recommendation["limitations"][0] in text
    document = Document(output)
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"] == [
        "一、关键发现", "二、文件分类", "三、推荐研究方向", "四、候选研究方向",
        "五、验证步骤与限制", "附件一 证据与来源", "附件二 研究规划原始记录"]
    body = text.split("附件一", 1)[0]
    assert "papers/" not in body and "C-last" not in body
    assert recommendation['counterevidence'][0]['text'] not in body
    assert document.sections[0].page_width.inches == 8.5
    assert document.styles["Normal"].font.name == "Microsoft YaHei"
    assert report == before


def test_explicit_counterevidence_outside_chain_keeps_independent_source_and_archive_locator(tmp_path):
    report = report_fixture()
    item = {"evidence_id": "C-archive", "source_path": "pack.zip::replication.pdf", "archive_source_path": "pack.zip",
            "page": 8, "text": "复核没有改善结果，不能据此宣称方法已通过验证。"}
    recommendation = report["recommended_research_direction"]
    recommendation["counterevidence"].append(item)
    recommendation["evidence_roles"]["support"].append("C-archive")
    _, text, _, _ = generated(tmp_path, report)
    assert "[C-archive] 来源：pack.zip::replication.pdf（归档来源：pack.zip），第 8 页；角色：反证线索" in text
    assert item["text"] in text
    assert text.count("[C-archive]") == 1


def test_role_only_counter_is_not_mislabeled_as_background_support(tmp_path):
    report = report_fixture()
    recommendation = report["recommended_research_direction"]
    recommendation["counterevidence"] = []
    _, text, _, _ = generated(tmp_path, report)
    counter = text.split("[C-last]", 1)[1].split("附件二", 1)[0]
    assert "背景支撑" not in counter
    assert "不能改善准确率|不适用于噪声标准差>0.4的样本。" in counter


def test_described_resources_still_require_confirmation_and_absent_counter_is_honest(tmp_path):
    report = report_fixture()
    recommendation = report["recommended_research_direction"]
    recommendation["quality_gate"]["resources"] = "material_described"
    recommendation["counterevidence"] = []
    recommendation["evidence_chain"] = recommendation["evidence_chain"][:-1]
    _, text, _, _ = generated(tmp_path, report)
    assert "资料已描述资源条件；仍需确认数据、代码、权限与运行条件的实际可用性。" in text
    assert "当前未提供可展示的反证线索，不表示不存在反例。" in text


def test_bounded_candidate_table_keeps_real_evidence_identity_counts_and_last_alternative(tmp_path):
    report = report_fixture()
    report["direction_candidates"] = [{"title": "候选%s" % i, "unique_evidence_count": 3, "independent_source_count": 2} for i in range(8)]
    output, text, _, _ = generated(tmp_path, report)
    table = Document(output).tables[0]
    assert table.rows[0].cells[-1].text == "证据/来源身份"
    assert len(table.rows) == 9
    assert table.rows[-1].cells[1].text == "候选7"
    assert table.rows[-1].cells[-1].text == "3/2"
    assert "来源身份数不等于独立实验数量" in text


@pytest.mark.parametrize("large_marker", ("report_kind", "scan_type", "package_id"))
def test_large_and_legacy_docx_keep_identical_body_and_styles(tmp_path, large_marker):
    legacy = report_fixture()
    legacy.pop("research_version")
    _, legacy_text, legacy_xml, legacy_styles = generated(tmp_path, legacy, name="legacy.docx")
    assert "为什么优先这个" not in legacy_text and "研究假设与状态" not in legacy_text
    current = report_fixture()
    scan = copy.deepcopy(SCAN)
    if large_marker == "report_kind": current["kind"] = "large"
    elif large_marker == "scan_type": scan["package_type"] = "large"
    else: scan["scan_id"] = "lp-fixture"
    _, _, xml, styles = generated(tmp_path, current, scan, name="large.docx")
    assert xml == legacy_xml
    assert styles == legacy_styles


def test_export_node_and_summary_download_code_remain_frozen():
    source = Path(__file__).resolve().parents[1]
    baseline = source.parents[1] / "model_decoupling_merged_20261011" / "source" / "services" / "exporter.py"
    if not baseline.is_file():
        pytest.skip("frozen 673-file checkout not present")
    before = ast.parse(baseline.read_text(encoding="utf-8-sig"))
    after = ast.parse((source / "services" / "exporter.py").read_text(encoding="utf-8-sig"))
    frozen_functions = {node.name: ast.dump(node) for node in before.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name != "create_report_docx"}
    current_functions = {node.name: ast.dump(node) for node in after.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name != "create_report_docx"}
    assert frozen_functions == current_functions


def test_iso_timestamp_is_chinese_in_body_and_original_is_retained_in_appendix(tmp_path):
    scan = dict(SCAN, scanned_at='2026-10-11T09:00:01Z')
    output, text, _, _ = generated(tmp_path, report_fixture(), scan)
    document=Document(output)
    paragraphs=[p.text for p in document.paragraphs]
    appendix=paragraphs.index('附件一 证据与来源')
    assert not re.search('[A-Za-z]', '\n'.join(paragraphs[:appendix]))
    assert '2026年10月11日 09:00:01（时区+00:00）' in text
    assert scan['scanned_at'] in '\n'.join(paragraphs[appendix:])
