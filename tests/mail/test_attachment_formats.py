"""Attachment execution routes, truthful states and complete mailbox locators."""
import io
import json
import mailbox
from dataclasses import replace
from email.message import EmailMessage
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

import pytest

from config import Config
from services.file_capabilities import (
    ARCHIVE_EXTENSIONS,
    IMAGE_EXTENSIONS,
    MAIL_SOURCE_EXTENSIONS,
    TEXT_EXTENSIONS,
)
from services.mail.extraction import _extract_attachment, extract_ocr
from services.mail.models import SourceRef
from services.mail.parsers.eml import parse_eml
from services.mail.policy import MailPolicy, MailPolicyError
from services.mail.source import MboxAdapter, NativeAdapter, open_locator


def attachment(payload, suffix, content_type='application/octet-stream'):
    return SimpleNamespace(raw_payload=payload, filename_raw='attachment' + suffix,
                           content_type=content_type, charset=None)


@pytest.mark.parametrize('suffix', sorted(TEXT_EXTENSIONS))
def test_text_formats_execute_without_a_text_mime(suffix):
    result = _extract_attachment(attachment('真实正文 19,000.00'.encode(), suffix))
    assert result['status'] == 'ready'
    assert '真实正文 19,000.00' in result['text']


@pytest.mark.parametrize('suffix', sorted(ARCHIVE_EXTENSIONS | MAIL_SOURCE_EXTENSIONS))
def test_nested_containers_do_not_recurse_or_decode_as_text(suffix):
    result = _extract_attachment(attachment(b'private nested bytes', suffix, 'text/plain'))
    assert result['status'] == 'restricted'
    assert result['text'] == ''


@pytest.mark.parametrize('suffix', sorted(IMAGE_EXTENSIONS))
def test_images_with_binary_mime_wait_for_ocr(suffix):
    from PIL import Image
    output = io.BytesIO()
    Image.new('RGB', (8, 8), 'white').save(output, format='PNG')
    result = _extract_attachment(attachment(output.getvalue(), suffix))
    assert result['status'] == 'ocr_pending'
    assert result['coverage_ratio'] == 0


def test_corrupt_image_does_not_claim_ocr_success():
    result = _extract_attachment(attachment(b'broken image', '.png'))
    assert result['status'] == 'failed'
    assert result['text'] == ''
    with patch('services.mail.extraction._image_metadata', side_effect=ImportError('Pillow')):
        missing = _extract_attachment(attachment(b'broken image', '.png'))
    assert missing['status'] == 'unsupported'
    assert 'Pillow' in missing['warnings'][0]


def test_empty_native_text_is_not_ready_and_read_budget_is_partial():
    empty = _extract_attachment(attachment(b'', '.json'))
    assert empty['status'] == 'partial'
    assert empty['coverage_ratio'] == 0
    limited = _extract_attachment(attachment(b'x' * 100, '.txt'), max_chars=8)
    assert limited['status'] == 'partial'
    assert len(limited['text']) <= 8
    assert limited['warnings']


@pytest.mark.parametrize('suffix', ['.pdf', '.docx', '.pptx', '.xlsx', '.xlsm', '.doc', '.ppt', '.rtf', '.png'])
def test_document_extension_wins_over_misleading_text_mime(suffix):
    result = _extract_attachment(attachment(b'corrupt document bytes', suffix, 'text/plain'))
    assert result['status'] != 'ready'
    assert result['parser'] != 'mail-text'


@pytest.mark.parametrize('suffix', ['.docx', '.pptx', '.xlsx', '.xlsm'])
def test_ooxml_preflight_rejects_high_ratio_before_parser(suffix):
    import zipfile
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as package:
        package.writestr('bomb.xml', b'x' * 500000)
    result = _extract_attachment(attachment(output.getvalue(), suffix))
    assert result['status'] == 'failed'
    assert result['error']['type'] == 'MailPolicyError'
    assert '压缩比' in result['error']['message']


def test_pptx_extracts_slides_tables_and_notes_with_location():
    from pptx import Presentation
    from pptx.util import Inches
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = '幻灯片报价 23,000'
    table = slide.shapes.add_table(1, 2, 0, Inches(1), Inches(3), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = '条款', '审批后生效'
    slide.notes_slide.notes_text_frame.text = '备注附加条件'
    output = io.BytesIO()
    document.save(output)
    result = _extract_attachment(attachment(output.getvalue(), '.pptx'))
    assert result['status'] == 'ready'
    for token in ['幻灯片报价 23,000', '条款', '审批后生效', '备注附加条件']:
        assert token in result['text']
    assert all(unit['page_number'] == 1 for unit in result['units'])
    assert {unit['document_part'] for unit in result['units']} == {'slide', 'notes'}
    misleading = _extract_attachment(attachment(output.getvalue(), '.pptx', 'application/pdf'))
    assert misleading['status'] == 'ready'
    assert misleading['parser'] == 'python-pptx'
    assert _extract_attachment(attachment(b'corrupt OOXML', '.pptx'))['status'] == 'failed'


@pytest.mark.parametrize('suffix,mime', [('.doc', 'application/msword'), ('.ppt', 'application/vnd.ms-powerpoint'), ('.rtf', 'text/rtf')])
def test_legacy_attachment_uses_bounded_helper_and_maps_incomplete_result(suffix, mime):
    calls = []

    def extract(path, **kwargs):
        calls.append((Path(path).suffix, Path(path).read_bytes(), kwargs))
        return {'text': '旧格式正文', 'parser': 'fixture', 'truncated': True,
                'metadata': {'content_complete': False}, 'warnings': ['bounded']}

    with patch('services.legacy_office.extract_legacy', side_effect=extract):
        result = _extract_attachment(attachment(b'legacy bytes', suffix, mime), max_chars=100)
    assert result['status'] == 'partial'
    assert result['coverage_ratio'] is None
    assert calls[0] == (suffix, b'legacy bytes', {'max_chars': 100, 'seconds': Config.MAIL_PARSE_TIMEOUT_SECONDS})
    with patch('services.legacy_office.extract_legacy', return_value={'text': '', 'warnings': ['缺少 striprtf 组件']}):
        failed = _extract_attachment(attachment(b'legacy bytes', suffix, mime))
    assert failed['status'] == 'failed'
    assert 'striprtf' in failed['error']['message']


def test_multiframe_image_ocr_reports_unprocessed_frames():
    from PIL import Image
    output = io.BytesIO()
    Image.new('RGB', (8, 8), 'white').save(output, format='TIFF', save_all=True,
        append_images=[Image.new('RGB', (8, 8), 'black')])

    def complete(command, **kwargs):
        Path(command[3]).with_name('ocr-result.json').write_text(json.dumps({
            'status': 'ready', 'text': '首帧文字', 'coverage_ratio': 1}), encoding='utf-8')
        return SimpleNamespace(returncode=0)

    with patch('services.mail.ocr.available', return_value=True), patch('subprocess.run', side_effect=complete):
        result = extract_ocr(attachment(output.getvalue(), '.tif'), object())
    assert result['status'] == 'partial'
    assert result['coverage_ratio'] == .5
    assert result['meta']['processed_image_frames'] == 1


@pytest.mark.parametrize('suffix', ['.mbox', '.mbx'])
def test_mailbox_import_keeps_each_real_message_locator(suffix):
    with TemporaryDirectory() as folder:
        path = Path(folder) / ('mailbox' + suffix)
        box = mailbox.mbox(str(path))
        for index in range(2):
            message = EmailMessage()
            message['Message-ID'] = f'<{index}@example.test>'
            message['Subject'] = f'message {index}'
            message.set_content(f'body {index}')
            box.add(message)
        box.close()
        policy = MailPolicy.from_config(Config)
        source = SourceRef('package', 'mbox', str(path), 'revision')
        adapter = MboxAdapter()
        assert adapter.supports(source)
        members = list(adapter.inventory(source, policy))
        assert len(members) == 2
        assert members[0].locator != members[1].locator
        for index, member in enumerate(members):
            locator = json.loads(json.dumps(member.locator))
            with open_locator(locator) as stream:
                parsed = parse_eml(stream.read())
            assert parsed.message_id_raw == [f'<{index}@example.test>']
        with pytest.raises(MailPolicyError):
            list(adapter.inventory(source, replace(policy, max_message_bytes=5)))
        with pytest.raises(MailPolicyError):
            list(adapter.inventory(source, replace(policy, max_total_uncompressed_bytes=members[0].byte_size)))


@pytest.mark.parametrize('name,suffix,dependency', [('msg', '.msg', 'extract-msg'), ('pst', '.pst', 'readpst')])
def test_missing_native_mail_dependency_is_explicit(name, suffix, dependency):
    adapter = NativeAdapter(name, {suffix}, dependency)
    with patch.object(adapter, 'available', return_value=False), pytest.raises(MailPolicyError, match=dependency):
        list(adapter.inventory(SourceRef('package', name, 'mail' + suffix, 'revision'), MailPolicy.from_config(Config)))


def test_native_cache_is_complete_only_after_valid_eml_output(tmp_path):
    adapter = NativeAdapter('msg', {'.msg'}, 'extract-msg')
    source = tmp_path / 'source.msg'
    source.write_bytes(b'native fixture')
    policy = MailPolicy.from_config(Config)

    def convert(command, **kwargs):
        cache = Path(command[-1])
        (cache / 'message.eml').write_bytes(b'Message-ID: <native@example.test>\r\n\r\nbody')
        return SimpleNamespace(returncode=0)

    with patch.object(Config, 'DATA_DIR', tmp_path), patch.object(adapter, 'available', return_value=True):
        with patch('subprocess.run', return_value=SimpleNamespace(returncode=0)), pytest.raises(MailPolicyError, match='EML'):
            list(adapter.inventory(SourceRef('package', 'msg', str(source), 'empty'), policy))
        assert not list(tmp_path.rglob('complete.json'))
        with patch('subprocess.run', side_effect=convert) as run:
            members = list(adapter.inventory(SourceRef('package', 'msg', str(source), 'valid'), policy))
            assert len(members) == 1
            assert len(list(tmp_path.rglob('complete.json'))) == 1
            assert len(list(adapter.inventory(SourceRef('package', 'msg', str(source), 'valid'), policy))) == 1
            assert run.call_count == 1


@pytest.mark.parametrize('suffix', ['.rtf', '.doc'])
def test_actual_legacy_worker_extracts_rtf_payload(tmp_path, suffix):
    # Word RTF disguised as .doc is handled by the worker's signature sniffing.
    pytest.importorskip('striprtf')
    payload = br'{\rtf1\ansi Quotation 23000.00\par Approved after review.}'
    result = _extract_attachment(attachment(payload, suffix))
    assert result['status'] == 'ready', result
    assert 'Quotation 23000.00' in result['text']
    assert 'Approved after review.' in result['text']
