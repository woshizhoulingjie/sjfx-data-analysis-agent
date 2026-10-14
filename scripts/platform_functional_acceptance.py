"""Real isolated SJFX HTTP/worker acceptance, reproducible on Windows and Linux.

Run with the application Python, e.g. python scripts/platform_functional_acceptance.py
--evidence /absolute/path/out. Only synthetic data is used. A new instance is
created for every run; model services are disabled unless explicit authorized
endpoints are supplied. --extended-model adds bounded representative ordinary,
large-package and mail model workflows. OCR quality and offline translation
need separate resource acceptance. Passing samples does not certify all domains.
"""
from __future__ import annotations

import argparse
import ast
from email.message import EmailMessage
import hashlib
import io
import json
import os
from pathlib import Path
import platform
import re
import secrets
import socket
import subprocess
import sys
import time
import traceback
import zipfile

ROOT = Path(__file__).resolve().parents[1]
MARKER = "SJFX_ACCEPTANCE_20261013"


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def inventory():
    routes = []
    for source in [ROOT / "app.py", *sorted((ROOT / "services").rglob("*.py")),
                   *sorted((ROOT / "sjfx_multiuser").rglob("*.py")),
                   *sorted((ROOT / "sjfx_protection").rglob("*.py"))]:
        tree = ast.parse(source.read_text(encoding="utf-8-sig"))
        for node in ast.walk(tree):
            if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                continue
            for deco in node.decorator_list:
                if not (isinstance(deco, ast.Call) and isinstance(deco.func, ast.Attribute)
                        and deco.func.attr in {"route", "get", "post", "put", "delete", "patch"}
                        and deco.args and isinstance(deco.args[0], ast.Constant)):
                    continue
                methods = [deco.func.attr.upper()] if deco.func.attr != "route" else ["GET"]
                for kw in deco.keywords:
                    if kw.arg == "methods":
                        methods = ast.literal_eval(kw.value)
                routes.append({"path": deco.args[0].value, "methods": methods,
                               "source": source.relative_to(ROOT).as_posix(),
                               "line": node.lineno, "handler": node.name})
    return {"routes": routes,
            "templates": [p.relative_to(ROOT).as_posix() for p in sorted((ROOT / "templates").rglob("*")) if p.is_file()],
            "tests": [p.relative_to(ROOT).as_posix() for p in sorted((ROOT / "tests").rglob("test*")) if p.is_file()],
            "workers": ["worker.py", "large_package_worker.py", "scripts/large_qa_worker.py"],
            "entrypoint": "sjfxctl.py"}


class Acceptance:
    def __init__(self, evidence, timeout, generation_url=None, embedding_url=None, model="qwen3.6:27b", extended_model=False, scenario="all"):
        import httpx
        self.httpx = httpx
        self.timeout = timeout
        self.generation_url, self.embedding_url, self.model = generation_url, embedding_url, model
        self.extended_model = extended_model
        self.scenario = scenario
        self.root = evidence.resolve() / (time.strftime("run_%Y%m%d_%H%M%S_") + secrets.token_hex(3))
        self.root.mkdir(parents=True)
        self.instance = self.root / "独立 实例"
        self.instance.mkdir()
        self.env_file = self.instance / "instance.env"
        self.state = self.instance / "state"
        self.output = self.instance / "outputs"
        self.corpus = self.instance / "合成材料"
        self.corpus.mkdir()
        self.results = []
        self.ids = {}
        self.clients = []
        self.started = False
        self.password = secrets.token_urlsafe(20)
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            self.port = sock.getsockname()[1]
        self.base = "http://127.0.0.1:" + str(self.port)
        # Do not inherit a production .env, token, endpoint or state path.
        self.env = {k: v for k, v in os.environ.items()
                    if not k.startswith(("SJFX_", "VLLM_", "OLLAMA_", "OPENAI_COMPATIBLE_", "MAIL_"))}
        self.env.update({"SJFX_ENV_FILE": str(self.env_file), "PYTHONUTF8": "1",
                         "PYTHONIOENCODING": "utf-8", "HF_HUB_OFFLINE": "1",
                         "TRANSFORMERS_OFFLINE": "1"})
        runtime_temp = self.instance / "runtime-temp"
        runtime_temp.mkdir()
        self.env.update({"TEMP": str(runtime_temp), "TMP": str(runtime_temp), "TMPDIR": str(runtime_temp), "SQLITE_TMPDIR": str(runtime_temp)})

    def cli(self, *args, check=True):
        proc = subprocess.run([sys.executable, "-X", "utf8", str(ROOT / "sjfxctl.py"),
                               "--env-file", str(self.env_file), *args],
                              cwd=ROOT, env=self.env, capture_output=True, text=True,
                              encoding="utf-8", errors="replace", timeout=self.timeout + 40)
        name = "cli_" + "_".join(args[:2]).replace("-", "_") + "_" + str(len(list(self.root.glob("cli_*"))))
        (self.root / (name + ".log")).write_text(proc.stdout + "\n" + proc.stderr, encoding="utf-8")
        if check and proc.returncode:
            raise AssertionError("CLI failed {}: {}".format(args, (proc.stdout + proc.stderr)[-3000:]))
        try:
            return json.loads(proc.stdout)
        except ValueError:
            return {"exit_code": proc.returncode, "stdout": proc.stdout[-4000:]}

    def client(self):
        client = self.httpx.Client(base_url=self.base, timeout=30, trust_env=False, follow_redirects=False)
        self.clients.append(client)
        return client

    def req(self, client, method, path, expected=(200,), **kwargs):
        response = client.request(method, path, **kwargs)
        if response.status_code not in expected:
            raise AssertionError("{} {} -> {}: {}".format(method, path, response.status_code, response.text[:2500]))
        if "application/json" in response.headers.get("content-type", ""):
            return response.json()
        return response

    def case(self, name, call):
        started = time.monotonic()
        print(name, flush=True)
        try:
            details = call()
            row = {"name": name, "status": "passed", "seconds": time.monotonic() - started, "details": details}
        except Exception as exc:
            row = {"name": name, "status": "failed", "seconds": time.monotonic() - started,
                   "error": str(exc), "traceback": traceback.format_exc()}
            print("FAILED: " + str(exc)[:1800] + (" [full details in acceptance.json]" if len(str(exc)) > 1800 else ""), flush=True)
        self.results.append(row)
        self.report()
        return row["status"] == "passed"

    def report(self):
        write_json(self.root / "acceptance.json", {
            "platform": platform.platform(), "python": sys.version, "source": str(ROOT),
            "instance": str(self.instance), "base_url": self.base,
            "model_scope": ("real endpoint " + self.generation_url if self.generation_url else "disabled; no inference quality claim"), "scenario": self.scenario, "ids": self.ids,
            "passed": sum(r["status"] == "passed" for r in self.results),
            "failed": sum(r["status"] == "failed" for r in self.results), "cases": self.results,
            "not_exercised": (["domain-wide model quality beyond synthetic representative samples"] if self.generation_url else
                              ["real model generation and QA quality"]) + ["real OCR image recognition",
                              "offline NLLB translation", "other operating systems",
                              "browser interaction and layout", "production-scale resource endurance"]})

    def fixtures(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        text = "合成验收材料 " + MARKER + "。2026年10月13日采购合计120元，共3项；仅用于隔离测试。"
        for name in ("采购说明.txt", "说明.md"):
            (self.corpus / name).write_text(text, encoding="utf-8")
        (self.corpus / "台账.csv").write_text("item,amount,marker\nA,40," + MARKER + "\nB,80," + MARKER + "\n", encoding="utf-8-sig")
        write_json(self.corpus / "资料.json", {"marker": MARKER, "amount": 120, "scope": "synthetic"})
        (self.corpus / "页面.html").write_text("<html><body><h1>合成资料</h1><p>" + text + "</p></body></html>", encoding="utf-8")
        doc = Document(); doc.add_heading("合成采购记录", 0); doc.add_paragraph(text); doc.save(self.corpus / "报告.docx")
        wb = Workbook(); wb.active.append(["marker", "amount"]); wb.active.append([MARKER, 120]); wb.save(self.corpus / "统计.xlsx")
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6]);
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.text = text
        deck.save(self.corpus / "演示.pptx")
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter(); page = writer.add_blank_page(612, 792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject(); content.set_data(("BT /F1 12 Tf 40 700 Td (" + MARKER + " amount 120 synthetic evidence) Tj ET").encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
        with (self.corpus / "文本.pdf").open("wb") as handle: writer.write(handle)
        with zipfile.ZipFile(self.corpus / "档案.zip", "w") as archive:
            archive.writestr("内部记录.txt", text.encode("utf-8"))
        self.mail = self.instance / "合成邮件"; self.mail.mkdir()
        for index in range(2):
            message = EmailMessage(); message["From"] = "alice@example.test"; message["To"] = "bob@example.test"
            message["Subject"] = "合成采购确认 " + MARKER; message["Message-ID"] = "<acceptance-{}@example.test>".format(index)
            message["Date"] = "Tue, 13 Oct 2026 10:00:00 +0800"
            if index: message["In-Reply-To"] = "<acceptance-0@example.test>"
            message.set_content(text + " 邮件编号=" + str(index))
            if not index: message.add_attachment((self.corpus / "台账.csv").read_bytes(), maintype="text", subtype="csv", filename="台账.csv")
            (self.mail / ("邮件" + str(index) + ".eml")).write_bytes(message.as_bytes())
        write_json(self.root / "fixture_hashes.json", {str(p.relative_to(self.instance)): hashlib.sha256(p.read_bytes()).hexdigest()
                                                       for folder in (self.corpus, self.mail) for p in folder.iterdir()})
        return {"ordinary_files": len(list(self.corpus.iterdir())), "mail_messages": 2, "marker": MARKER}

    def setup(self):
        self.cli("setup", "--data-dir", str(self.state), "--output-dir", str(self.output),
                 "--models-dir", str(self.instance / "models"), "--backend", "openai_compatible", "--port", str(self.port))
        overrides = ('SCAN_ALLOWED_ROOTS="' + str(self.instance) + '"\n' +
                         "ENABLE_VLLM=0\nENABLE_SHARED_OLLAMA=0\nENABLE_OPENAI_COMPATIBLE=0\n" +
                         "OPENAI_COMPATIBLE_BASE_URL=http://127.0.0.1:9/v1\n" +
                         "VLLM_BASE_URL=http://127.0.0.1:9/v1\nOLLAMA_BASE_URL=http://127.0.0.1:9\n" +
                         "SJFX_EMBED_BACKEND=disabled\nSJFX_EMBED_BASE_URL=http://127.0.0.1:9/v1\n" +
                         "ENABLE_TRANSLATION=0\nENABLE_IMPORT_TRANSLATION=0\nAUTO_TRANSLATE_PACKAGES=0\n" +
                         "MAIL_ENABLE_VECTOR_INDEX=0\nMAIL_ENABLE_FACT_EXTRACTION=0\nMAIL_ENABLE_AGENT_MINING=0\n" +
                         "JOB_POLL_SECONDS=0.2\nSOURCE_STABILITY_SECONDS=0\n")
        replacement_keys = {line.split("=", 1)[0] for line in overrides.splitlines() if "=" in line}
        if self.generation_url:
            overrides += "LLM_BACKEND=vllm\nENABLE_VLLM=1\nVLLM_BASE_URL=" + self.generation_url + "\nVLLM_MODEL=" + self.model + "\nLLM_CONTEXT_TOKENS=32768\n"
            if self.extended_model:
                overrides += ("LARGE_QA_DEMAND_ONLY=0\nLARGE_QA_IDLE_SLEEP=1\nLARGE_QA_EXTRACT_PAUSE=0\n"
                              "LARGE_QA_SYNC_INTERVAL=2\nLARGE_QA_REFINE_TOPICS=0\n"
                              "MAIL_RESEARCH_MAX_MODEL_CALLS=16\nMAIL_RESEARCH_MAX_TOOL_CALLS=30\n"
                              "MAIL_RESEARCH_MAX_TOKENS=60000\nMAIL_RESEARCH_TIMEOUT_SECONDS=" + str(self.timeout) + "\n")
            if self.embedding_url:
                overrides += "SJFX_EMBED_BACKEND=openai_compatible\nSJFX_EMBED_BASE_URL=" + self.embedding_url + "\nSJFX_EMBED_MODEL=multilingual-e5-small\n"
            # Keep one final value per key.
            values = dict(line.split("=", 1) for line in overrides.splitlines() if "=" in line)
            overrides = "\n".join(key + "=" + value for key, value in values.items()) + "\n"
            replacement_keys = set(values)
        # Config's env loader preserves the first assignment. Replace setup
        # entries, never append conflicting duplicate keys.
        original = self.env_file.read_text(encoding="utf-8")
        retained = [line for line in original.splitlines() if line.split("=", 1)[0] not in replacement_keys]
        self.env_file.write_text("\n".join(retained) + "\n" + overrides, encoding="utf-8")
        return {"env_file": str(self.env_file), "isolated": True}

    def start(self):
        self.started = True  # cleanup even after a partially failed start
        status = self.cli("start", "--timeout", str(self.timeout))
        assert status.get("ready"), status
        assert set(status["roles_alive"]) == {"web", "worker", "large-worker", "qa-worker"}, status
        assert all(status["roles_alive"].values()), status
        return {"roles_alive": status["roles_alive"], "ready": status["ready"], "run_id": status["run_id"]}

    def auth(self):
        self.anonymous = self.client(); self.admin = self.client(); self.user = self.client(); self.other = self.client()
        self.req(self.anonymous, "GET", "/api/auth/me", expected=(401,))
        log = "\n".join(path.read_text(encoding="utf-8", errors="replace")
                        for path in (self.state / "logs").glob("*.log"))
        match = re.search(r"初始密码：([^（\s]+)", log)
        assert match, "first admin initial password missing from isolated app log"
        initial = match.group(1)
        self.req(self.admin, "POST", "/login", expected=(302, 303), data={"username": "admin", "password": initial})
        first = self.req(self.admin, "GET", "/api/auth/me")["user"]
        assert first["role"] == "admin" and first["must_change_pw"]
        self.req(self.admin, "GET", "/api/admin/model-settings", expected=(403,))
        self.req(self.admin, "POST", "/api/auth/change-password", data={"current_password": initial, "new_password": self.password, "confirm_password": self.password})
        self.req(self.admin, "POST", "/login", expected=(302, 303), data={"username": "admin", "password": self.password})
        for name, client in (("acceptance_user", self.user), ("acceptance_other", self.other)):
            self.req(client, "POST", "/register", expected=(302, 303), data={"username": name, "password": self.password, "confirm": self.password, "role": "admin"})
            profile = self.req(client, "GET", "/api/auth/me")["user"]
            assert profile["role"] == "user", profile
            self.req(client, "GET", "/api/admin/model-settings", expected=(403,))
        return {"first_admin_must_change_password": True, "user_role_escalation_rejected": True,
                "user_owner": self.req(self.user, "GET", "/api/auth/me")["user"]["owner_id"]}

    def pages(self):
        checked = []
        for path in ("/", "/login", "/register", "/admin", "/mail-research", "/protection", "/protection/admin"):
            client = self.anonymous if path in {"/login", "/register"} else self.admin
            body = self.req(client, "GET", path)
            assert body.text and "text/html" in body.headers.get("content-type", ""), path
            checked.append(path)
        return {"HTTP_shells": checked, "browser_visual_QA": "not exercised"}

    def diagnostics(self):
        status = self.req(self.user, "GET", "/api/status")
        assert status["runtime_supported"] and status["state_dir"] == str(self.state), status
        assert status["model_generation_enabled"] is bool(self.generation_url), status
        caps = self.req(self.user, "GET", "/api/file-capabilities")
        assert caps.get("ok") and len(json.dumps(caps)) > 300
        settings = self.req(self.admin, "GET", "/api/admin/model-settings")
        self.req(self.user, "POST", "/api/admin/model-test", expected=(403,), json={"role": "generation", "timeout_seconds": 1})
        return {"status": status, "capabilities": caps, "settings": settings, "doctor": self.cli("doctor", check=False)}

    def wait_job(self, job_id, *, expect="completed"):
        end = time.monotonic() + self.timeout
        while time.monotonic() < end:
            item = self.req(self.user, "GET", "/api/jobs/" + job_id)["job"]
            if item["status"] in {"completed", "failed", "cancelled"}:
                write_json(self.root / ("job_" + job_id + ".json"), item)
                assert item["status"] == expect, item
                return item
            time.sleep(.3)
        raise TimeoutError("job did not reach terminal status: " + job_id)

    def wait_package_jobs(self, package):
        """Wait for import follow-up work before taking a fixed mail snapshot."""
        end = time.monotonic() + self.timeout
        stable = 0
        while time.monotonic() < end:
            jobs = self.req(self.user, "GET", "/api/jobs", params={"status": "active", "limit": 500})["jobs"]
            related = [j for j in jobs if j.get("scan_id") == package]
            stable = stable + 1 if not related else 0
            if stable >= 3: return {"stable_checks": stable}
            time.sleep(.5)
        raise TimeoutError("package follow-up tasks remain active: " + package)

    def scan(self):
        result = self.req(self.user, "POST", "/api/scan", expected=(202,), json={"path": str(self.corpus), "parse_mode": "fast"})
        job = self.wait_job(result["job_id"])
        scan_id = (job.get("result") or {}).get("scan_id") or result["job_id"]
        self.ids["scan"] = scan_id
        scan = self.req(self.user, "GET", "/api/scan/" + scan_id)
        assert scan["scan"]["file_count"] >= 10, scan
        self.req(self.other, "GET", "/api/scan/" + scan_id, expected=(404,))
        files = self.req(self.user, "GET", "/api/scan/" + scan_id + "/files")
        selection = self.req(self.user, "GET", "/api/scan/" + scan_id + "/selection")
        sources = self.req(self.user, "GET", "/api/data-sources")
        assert any(item.get("scan_id") == scan_id or item.get("id") == scan_id for item in sources["items"]), sources
        reused = self.req(self.user, "POST", "/api/data-sources/select", json={"scan_id": scan_id})
        assert reused["created_job"] is False
        return {"job": job, "scan": scan, "files": files, "selection": selection, "historical_reuse": reused}

    def upload(self):
        data = (MARKER + " private upload content").encode()
        sha = hashlib.sha256(data).hexdigest()
        result = self.req(self.user, "POST", "/api/uploads", json={"files": [{"rel_path": "私有材料.txt", "size": len(data), "sha256": sha}]})
        sid = result["session_id"]
        self.req(self.user, "PUT", "/api/uploads/" + sid + "/chunks/0", params={"chunk_index": 0, "sha256": sha}, headers={"Content-Type": "application/octet-stream"}, content=data)
        done = self.req(self.user, "POST", "/api/uploads/" + sid + "/finalize", json={})
        assert (Path(done["dest_dir"]) / "私有材料.txt").read_bytes() == data
        self.req(self.other, "POST", "/api/scan", expected=(403,), json={"path": done["dest_dir"]})
        accepted = self.req(self.user, "POST", "/api/scan", expected=(202,), json={"path": done["dest_dir"]})
        job = self.wait_job(accepted["job_id"])
        self.ids["uploaded_scan"] = accepted["job_id"]
        return {"finalized": done, "job": job, "foreign_owner_denied": True, "sha256": sha}

    def raw_and_search(self):
        scan = self.ids["scan"]
        raw = self.req(self.user, "GET", "/api/document/" + scan + "/raw", params={"path": "采购说明.txt"})
        assert MARKER.encode() in raw.content
        search = self.req(self.user, "POST", "/api/search/files", json={"scan_id": scan, "query": MARKER})
        assert search["matched_file_count"] == 0 and search["coverage"]["complete"] is False and search.get("warnings"), search
        conversation = self.req(self.user, "POST", "/api/conversations", expected=(201,), json={"scan_id": scan})
        self.req(self.user, "POST", "/api/conversation/" + conversation["session"]["session_id"] + "/turns", expected=(409,), json={"scan_id": scan, "question": "采购合计多少？"})
        return {"raw_sha256": hashlib.sha256(raw.content).hexdigest(), "search": search,
                "search_scope": "inventory awaiting selection; verified explicit incomplete coverage and QA 409 guard; positive retrieval not certified"}

    def export(self):
        response = self.req(self.user, "POST", "/api/export", expected=(202,), json={"scan_id": self.ids["scan"], "path": ".", "kind": "directory", "task_topic": "合成采购资料交接"})
        job = self.wait_job(response["job_id"])
        name = job["result"]["file_name"]
        ticket = self.req(self.user, "POST", "/api/download-ticket", json={"filename": name})
        raw = self.req(self.user, "GET", ticket["download_url"])
        assert raw.content.startswith(b"PK") and len(raw.content) > 200
        with zipfile.ZipFile(io.BytesIO(raw.content)) as archive:
            names = archive.namelist()
            # The exporter deduplicates identical TXT/MD fixtures. Verify a
            # unique CSV source rather than requiring both duplicate names.
            assert any(name.endswith("台账.csv") for name in names), names
            match = next(name for name in names if name.endswith("台账.csv"))
            assert archive.read(match) == (self.corpus / "台账.csv").read_bytes()
            assert any(name.endswith("导出清单.json") for name in names), names
        (self.root / "ordinary_export.zip").write_bytes(raw.content)
        self.req(self.user, "GET", ticket["download_url"], expected=(401,))
        self.req(self.other, "POST", "/api/download-ticket", expected=(404,), json={"filename": name})
        self.req(self.other, "GET", "/outputs/" + name, expected=(404,))
        listed = self.req(self.user, "GET", "/api/exports")
        assert listed["total"] >= 1
        self.ids["export"] = name
        return {"job": job, "archive_members": names, "sha256": hashlib.sha256(raw.content).hexdigest(),
                "one_use_ticket": True, "foreign_owner_denied": True}

    def mail_import(self):
        result = self.req(self.user, "POST", "/api/mail/packages", expected=(200, 202), json={"source": {"path": str(self.mail), "kind": "directory"}, "display_name": "合成邮件验收"})
        package = result["package"]["package_id"]
        self.ids["mail"] = package
        job = self.wait_job(result["job_id"])
        detail = self.req(self.user, "GET", "/api/mail/packages/" + package)
        messages = self.req(self.user, "GET", "/api/mail/packages/" + package + "/messages")
        assert messages["messages"]["total"] == 2 and len(messages["messages"]["items"]) == 2, messages
        self.req(self.other, "GET", "/api/mail/packages/" + package, expected=(404,))
        key = messages["messages"]["items"][0]["message_key"]
        preview = self.req(self.user, "GET", "/api/mail/messages/" + key + "/preview", params={"package_id": package})
        assert MARKER in json.dumps(preview, ensure_ascii=False), preview
        threads = self.req(self.user, "GET", "/api/mail/packages/" + package + "/threads")
        people = self.req(self.user, "GET", "/api/mail/packages/" + package + "/people")
        foundation = self.req(self.user, "GET", "/api/mail/packages/" + package + "/foundation")
        return {"job": job, "detail": detail, "messages": messages, "preview": preview, "threads": threads, "people": people, "foundation": foundation}

    def large(self):
        result = self.req(self.user, "POST", "/api/large-packages", expected=(202,), json={"path": str(self.corpus)})
        package = result["package_id"]; self.ids["large"] = package
        end = time.monotonic() + self.timeout
        detail = {}
        while time.monotonic() < end:
            detail = self.req(self.user, "GET", "/api/large-packages/" + package)
            item = detail["large_package"]
            if int(item.get("counts", {}).get("total_files") or 0) >= 10:
                break
            time.sleep(.5)
        assert int(detail["large_package"].get("counts", {}).get("total_files") or 0) >= 10, detail
        self.req(self.other, "GET", "/api/large-packages/" + package, expected=(404,))
        files = self.req(self.user, "GET", "/api/large-packages/" + package + "/files")
        workflow = self.req(self.user, "GET", "/api/large-packages/" + package + "/workflow")
        return {"detail": detail, "files": files, "workflow": workflow,
                "scope": "real inventory only; model directory and QA untested"}

    def pause_resume(self):
        scan = self.ids["scan"]
        pause = self.req(self.user, "POST", "/api/package-processing/" + scan + "/pause", json={"reason": "合成验收检查点"})
        state = self.req(self.user, "GET", "/api/package-processing/" + scan)
        assert state["processing"]["state"] == "paused", state
        resume = self.req(self.user, "POST", "/api/package-processing/" + scan + "/resume", expected=(200, 202), json={})
        assert resume["accepted"] and resume["job_id"], resume
        cancel = self.req(self.user, "POST", "/api/jobs/" + resume["job_id"] + "/cancel", json={})
        cancelled = self.wait_job(resume["job_id"], expect="cancelled")
        return {"pause": pause, "state": state, "resume": resume, "cancel": cancel, "cancelled_job": cancelled,
                "scope": "persisted package controls and real queued/running job cancellation; blocked-I/O cancellation tested separately"}

    def restart(self):
        stop = self.cli("stop", "--timeout", "30")
        assert stop.get("running") is False, stop
        start = self.start()
        profile = self.req(self.user, "GET", "/api/auth/me")
        scan = self.req(self.user, "GET", "/api/scan/" + self.ids["scan"])
        assert scan["scan"]["file_count"] >= 10
        mail = self.req(self.user, "GET", "/api/mail/packages/" + self.ids["mail"]) if "mail" in self.ids else {}
        return {"start": start, "existing_session": profile, "scan_file_count": scan["scan"]["file_count"], "mail": mail}

    def model_sample(self):
        sample = self.instance / "小样本真模型"
        sample.mkdir()
        (sample / "采购凭证.txt").write_text("合成采购凭证。采购合计120元，共3项。" + MARKER + "。这是验收数据，不是实际业务交易。", encoding="utf-8")
        accepted = self.req(self.user, "POST", "/api/scan", expected=(202,), json={"path": str(sample), "parse_mode": "fast"})
        job = self.wait_job(accepted["job_id"])
        scan = (job.get("result") or {}).get("scan_id") or accepted["job_id"]
        self.ids["model_scan"] = scan
        selection = self.req(self.user, "GET", "/api/scan/" + scan + "/selection")
        confirm = self.req(self.user, "POST", "/api/scan/" + scan + "/selection/confirm", expected=(200, 202), json={"included_paths": ["采购凭证.txt"], "excluded_paths": [], "version": selection["selection"]["version"]})
        write_json(self.root / "model_sample_confirm.json", confirm)
        parsed = self.wait_job(confirm["job_id"])
        document = self.req(self.user, "GET", "/api/document/" + scan, params={"path": "采购凭证.txt"})
        assert MARKER in json.dumps(document, ensure_ascii=False), document
        write_json(self.root / "model_sample_document.json", document)
        search = self.req(self.user, "POST", "/api/search/files", json={"scan_id": scan, "query": "采购"})
        assert int(search["matched_file_count"]) >= 1 and search["evidence_count"] >= 1, search
        write_json(self.root / "model_sample_search.json", search)
        conversation = self.req(self.user, "POST", "/api/conversations", expected=(201,), json={"scan_id": scan})
        session = conversation["session"]["session_id"]
        turn = self.req(self.user, "POST", "/api/conversation/" + session + "/turns", expected=(200, 202), json={"scan_id": scan, "question": "根据采购凭证，采购合计多少元？请引用原文证据。"})
        write_json(self.root / "model_sample_turn_created.json", turn)
        qa_job = self.wait_job(turn["job_id"])
        turn_id = turn["turn"]["id"]
        final = self.req(self.user, "GET", "/api/turns/" + turn_id)
        write_json(self.root / "model_sample_qa.json", final)
        answer = final["turn"].get("result") or {}
        # Never search the whole JSON: tools can contain the expected number
        # even when the user-facing answer refuses or has zero citations.
        assert "120" in str(answer.get("answer") or ""), answer
        assert answer.get("citations") and answer.get("evidence_status") != "insufficient", answer
        assert answer.get("task_status") != "not_fulfilled", answer
        assert any(c.get("source_path") == "采购凭证.txt" and "120" in c.get("original_text", "")
                   and c["original_text"] in (sample / "采购凭证.txt").read_text(encoding="utf-8")
                   for c in answer["citations"]), answer
        assert final["turn"]["status"] == "completed", final
        self.req(self.other, "GET", "/api/turns/" + turn_id, expected=(404,))
        return {"scope": "one synthetic TXT; application on " + platform.system() + " connected to explicit authorized endpoint",
                "selected_parse_job": parsed, "search": search, "qa_job": qa_job, "turn": final,
                "quality_check": "expected numeric fact 120 and persisted evidence; other domain quality not certified"}

    def download_artifact(self, name, evidence_name):
        ticket = self.req(self.user, "POST", "/api/download-ticket", json={"filename": name})
        raw = self.req(self.user, "GET", ticket["download_url"])
        assert raw.content, "empty artifact: " + name
        (self.root / evidence_name).write_bytes(raw.content)
        self.req(self.other, "POST", "/api/download-ticket", expected=(404,), json={"filename": name})
        return raw.content

    def extended_ordinary(self):
        """Two related originals, actual file/folder model calls and a real DOCX."""
        from docx import Document
        sample = self.instance / "同类采购样本"; sample.mkdir()
        texts = {"采购凭证.txt": "文件编号：ACCEPT-2026-001\n日期：2026-10-13\n事项编号：SYNTHETIC-120\n主题：合成采购凭证\n采购合计120元，共3项。采购目的为测试泵密封圈。" + MARKER,
                 "采购复核.txt": "文件编号：ACCEPT-2026-002\n日期：2026-10-13\n事项编号：SYNTHETIC-120\n主题：合成采购复核\n引用文件ACCEPT-2026-001，该凭证所记采购合计120元，共3项；尚未实际付款，不能推断支付完成。" + MARKER}
        for name, text in texts.items(): (sample / name).write_text(text, encoding="utf-8")
        accepted = self.req(self.user, "POST", "/api/scan", expected=(202,), json={"path": str(sample), "parse_mode": "fast"})
        scanned = self.wait_job(accepted["job_id"])
        scan = (scanned.get("result") or {}).get("scan_id") or accepted["job_id"]
        self.ids["extended_ordinary"] = scan
        selection = self.req(self.user, "GET", "/api/scan/" + scan + "/selection")["selection"]
        confirm = self.req(self.user, "POST", "/api/scan/" + scan + "/selection/confirm", expected=(200, 202),
                           json={"included_paths": list(texts), "excluded_paths": [], "version": selection["version"]})
        self.wait_job(confirm["job_id"])
        homogeneous = self.req(self.user, "POST", "/api/homogeneous-analysis/" + scan, expected=(202,), json={})
        self.wait_job(homogeneous["job_id"])
        homogeneous = self.req(self.user, "GET", "/api/homogeneous-analysis/" + scan)
        write_json(self.root / "extended_ordinary_homogeneous.json", homogeneous)
        assert homogeneous["available"] and homogeneous["analysis"]["summary"]["eligible"], homogeneous
        assert {row["path"] for row in homogeneous["analysis"]["records"]["items"]} == set(texts), homogeneous
        assert {row["fields"]["document_number"] for row in homogeneous["analysis"]["records"]["items"]} == {
            "ACCEPT-2026-001", "ACCEPT-2026-002"}, homogeneous
        self.req(self.other, "GET", "/api/homogeneous-analysis/" + scan, expected=(400,))
        summaries = {}
        for path in texts:
            requested = self.req(self.user, "POST", "/api/summary", expected=(200, 202),
                                 json={"scan_id": scan, "path": path, "kind": "file", "force": True})
            if requested.get("job_id"): self.wait_job(requested["job_id"])
            final = self.req(self.user, "POST", "/api/summary", expected=(200,),
                             json={"scan_id": scan, "path": path, "kind": "file"})
            summary = final["summary"]
            assert "120" in json.dumps(summary, ensure_ascii=False), summary
            assert (summary.get("parser_info") or {}).get("local_model"), summary
            summaries[path] = summary
        combined = self.req(self.user, "POST", "/api/summary", expected=(200, 202),
                            json={"scan_id": scan, "path": ".", "kind": "directory", "force": True})
        if combined.get("job_id"): self.wait_job(combined["job_id"])
        combined = self.req(self.user, "POST", "/api/summary", expected=(200,),
                            json={"scan_id": scan, "path": ".", "kind": "directory"})
        assert (combined["summary"].get("parser_info") or {}).get("local_model"), combined
        report = self.req(self.user, "POST", "/api/report", expected=(202,), json={"scan_id": scan})
        job = self.wait_job(report["job_id"])
        overview = job["result"]["overview"]
        assert set(overview["report"]["report_scope"]["selected_paths"]) == set(texts), overview
        assert overview["report"]["report_scope"]["selected_bytes"] == sum((sample / path).stat().st_size for path in texts), overview
        raw = self.download_artifact(overview["file_name"], "representative_ordinary_report.docx")
        doc = Document(io.BytesIO(raw))
        paragraphs = "\n".join(p.text for p in doc.paragraphs)
        assert "文件分类" in paragraphs and "证据与来源" in paragraphs and "验证步骤与限制" in paragraphs, paragraphs
        assert len(paragraphs) > 400, paragraphs
        write_json(self.root / "extended_ordinary_summaries.json", {"file": summaries, "same_scope_directory": combined})
        write_json(self.root / "extended_ordinary_report.json", overview)
        return {"scan_id": scan, "file_summaries": list(summaries), "same_scope_directory_model": True,
                "report_sha256": hashlib.sha256(raw).hexdigest(), "report_characters": len(paragraphs),
                "homogeneous_eligible": True, "homogeneous_scope": "deterministic labelled-field analysis, independent of the model",
                "scope": "two related synthetic TXT originals; actual file/folder model calls and dedicated homogeneous analysis"}

    def extended_large(self):
        sample = self.instance / "大包采购样本"; sample.mkdir()
        texts = {"采购凭证.txt": "合成采购凭证，采购合计120元，共3项。这是测试记录。" + MARKER,
                 "风险说明.txt": "合成采购风险说明：以上采购尚未付款，须核对凭证金额120元，不能认为付款已经完成。" + MARKER}
        for name, text in texts.items(): (sample / name).write_text(text, encoding="utf-8")
        created = self.req(self.user, "POST", "/api/large-packages", expected=(202,), json={"path": str(sample)})
        package = created["package_id"]; self.ids["extended_large"] = package
        end = time.monotonic() + self.timeout
        while time.monotonic() < end:
            files = self.req(self.user, "GET", "/api/large-packages/" + package + "/files")
            detail = self.req(self.user, "GET", "/api/large-packages/" + package)
            quick = detail["large_package"].get("counts", {}).get("quick", {})
            if quick.get("completed", 0) == 2: break
            time.sleep(.5)
        assert quick.get("completed", 0) == 2, detail
        self.req(self.other, "GET", "/api/large-packages/" + package, expected=(404,))
        end = time.monotonic() + self.timeout
        state = {}
        while time.monotonic() < end:
            status_path = self.state / "large_qa" / "worker_status.json"
            if status_path.exists():
                state = json.loads(status_path.read_text(encoding="utf-8"))
                if state.get("packages", {}).get(package, {}).get("parse", {}).get("complete") == 2: break
            time.sleep(.5)
        write_json(self.root / "extended_large_original_index.json", state)
        assert state.get("packages", {}).get(package, {}).get("parse", {}).get("complete") == 2, state
        conversation = self.req(self.user, "POST", "/api/conversations", expected=(201,),
                                json={"scan_id": package, "package_type": "large"})
        session = conversation["session"]["session_id"]
        created_turn = self.req(self.user, "POST", "/api/conversation/" + session + "/turns", expected=(200, 202),
                                json={"scan_id": package, "package_type": "large", "question": "根据采购凭证，采购合计多少元？请引用原文证据。"})
        self.wait_job(created_turn["job_id"])
        final = self.req(self.user, "GET", "/api/turns/" + created_turn["turn"]["id"])
        write_json(self.root / "extended_large_qa.json", final)
        result = final["turn"]["result"]
        assert final["turn"]["status"] == "completed" and "120" in result.get("answer", ""), final
        assert result.get("citations") and result.get("evidence_status") != "insufficient", result
        assert any(c.get("source_path") in texts and "120" in c.get("original_text", "")
                   and c["original_text"] in texts[c["source_path"]] for c in result["citations"]), result
        self.req(self.other, "GET", "/api/turns/" + created_turn["turn"]["id"], expected=(404,))
        return {"package_id": package, "files": files, "qa": result, "scope": "two synthetic large-package originals, real QA and exact source quote"}

    def extended_mail_screen(self):
        package = self.ids["mail"]
        self.wait_package_jobs(package)
        created = self.req(self.user, "POST", "/api/mail/packages/" + package + "/intent-analyses", expected=(202,),
                           json={"kind": "screen", "question": "查找采购合计120元且说明为合成验收材料的邮件。",
                                 "scope": {"include_explicit_attachments": False}, "model_call_budget": 6, "max_tokens": 15000})
        self.wait_job(created["job_id"])
        final = self.req(self.user, "GET", "/api/mail/intent-analyses/" + created["turn_id"])
        write_json(self.root / "extended_mail_screen.json", final)
        assert final["status"] == "completed" and final["progress"]["complete"], final
        assert final["total"] == 2 and len(final["batches"]) == 2, final
        assert all(row["status"] == "matched" and row.get("evidence") for row in final["batches"]), final
        assert all(any("120" in e.get("quote_text", "") for e in row["evidence"]) for row in final["batches"]), final
        self.ids["mail_screen_result"] = final["progress"]["result_set_id"]
        self.req(self.other, "GET", "/api/mail/intent-analyses/" + created["turn_id"], expected=(404,))
        return final

    def extended_mail_research(self):
        package = self.ids["mail"]
        session = self.req(self.user, "POST", "/api/mail/research-sessions", expected=(201,),
                           json={"package_id": package, "title": "合成采购金额研究", "scope": {"include_explicit_attachments": False}})["session"]
        created = self.req(self.user, "POST", "/api/mail/research-sessions/" + session["session_id"] + "/turns", expected=(202,),
                           json={"question": "根据邮件原文，采购合计多少元？请引用直接原文证据，说明是合成验收数据。",
                                 "budget": {"max_model_calls": 16, "max_tool_calls": 30, "max_tokens": 60000,
                                            "max_seconds": self.timeout, "max_gap_rounds": 1, "max_gap_candidates": 4}})
        self.wait_job(created["job_id"])
        turn = self.req(self.user, "GET", "/api/mail/research-turns/" + created["turn"]["turn_id"])["turn"]
        write_json(self.root / "extended_mail_research.json", turn)
        final = turn["final"]
        assert turn["status"] == "completed" and "120" in final.get("answer", ""), turn
        assert final.get("evidence") and any("120" in e.get("quote_text", "") for e in final["evidence"]), final
        assert any(c.get("semantic_status") == "supported" and "120" in str(c.get("claim_text", ""))
                   for c in final.get("claims", [])), final
        assert final.get("result_set_id"), final
        self.ids["mail_research_result"] = final["result_set_id"]
        self.req(self.other, "GET", "/api/mail/research-turns/" + turn["turn_id"], expected=(404,))
        return {"turn": turn, "scope": "small retrieval research, numeric fact and strictly supported original evidence"}

    def extended_mail_exports(self):
        import csv
        from openpyxl import load_workbook
        result_id = self.ids.get("mail_research_result") or self.ids["mail_screen_result"]
        result = self.req(self.user, "GET", "/api/mail/result-sets/" + result_id)["result_set"]
        expected_keys = {item["item_id"] for item in result["items"] if item["item_type"] == "message"}
        assert expected_keys, result
        outputs = {}
        for format_name in ("csv", "html", "xlsx"):
            created = self.req(self.user, "POST", "/api/mail/exports", expected=(200, 202),
                               json={"result_set_id": result_id, "format": format_name})
            job = self.wait_job(created["job_id"])
            raw = self.download_artifact(job["result"]["filename"], "representative_mail_export." + format_name)
            if format_name == "csv":
                rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
                assert {row["message_key"] for row in rows} == expected_keys, rows
                assert all(MARKER in row["subject"] for row in rows), rows
            elif format_name == "html":
                text = raw.decode("utf-8")
                assert "<html" in text and MARKER in text and all(key in text for key in expected_keys), text
            else:
                workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                try:
                    rows = list(workbook["邮件研究结果"].values)
                    assert {row[0] for row in rows[1:]} == expected_keys, rows
                    assert all(MARKER in row[3] for row in rows[1:]), rows
                    assert "覆盖说明" in workbook.sheetnames and "中文研究结论" in workbook.sheetnames
                finally: workbook.close()
            outputs[format_name] = {"bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}
        self.req(self.other, "GET", "/api/mail/result-sets/" + result_id, expected=(404,))
        return {"result_set_id": result_id, "message_count": len(expected_keys), "artifacts": outputs}

    def run(self):
        write_json(self.root / "source_inventory.json", inventory())
        try:
            if not self.case("synthetic_format_fixtures", self.fixtures): return
            if not self.case("fresh_isolated_setup", self.setup): return
            if not self.case("four_process_roles_start", self.start): return
            if not self.case("first_admin_password_and_user_roles", self.auth): return
            self.case("page_http_shells", self.pages)
            self.case("runtime_model_and_file_diagnostics", self.diagnostics)
            if self.generation_url:
                if self.scenario == "all":
                    self.case("real_model_selected_import_positive_search_and_qa", self.model_sample)
                if self.extended_model:
                    if self.scenario in ("all", "ordinary"):
                        self.case("real_model_ordinary_file_and_same_scope_analysis_docx", self.extended_ordinary)
                    if self.scenario in ("all", "large-qa"):
                        self.case("real_model_large_package_qa_original_citation", self.extended_large)
                    if self.scenario == "all" and self.case("mail_real_import_previews_threads_people", self.mail_import):
                        self.case("real_model_mail_full_text_screen", self.extended_mail_screen)
                        self.case("real_model_mail_retrieval_research", self.extended_mail_research)
                        if self.ids.get("mail_research_result") or self.ids.get("mail_screen_result"):
                            self.case("mail_worker_csv_html_xlsx_exports", self.extended_mail_exports)
                return
            self.case("ordinary_real_scan_and_historical_reuse", self.scan)
            self.case("private_chunk_upload_owner_isolation", self.upload)
            if "scan" in self.ids:
                self.case("ordinary_raw_content_and_unindexed_search_guard", self.raw_and_search)
                self.case("ordinary_export_bytes_manifest_and_permissions", self.export)
                self.case("package_pause_resume_checkpoints", self.pause_resume)
            self.case("mail_real_import_previews_threads_people", self.mail_import)
            self.case("large_real_inventory_and_owner_isolation", self.large)
            if "scan" in self.ids:
                self.case("restart_preserves_sessions_and_business_data", self.restart)
        finally:
            if self.started:
                self.case("clean_four_role_shutdown", lambda: self.cli("stop", "--timeout", "30"))
            for client in self.clients: client.close()
            self.report()
            print("Evidence: " + str(self.root), flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--evidence", type=Path, required=True)
    parser.add_argument("--timeout", type=int, default=180)
    parser.add_argument("--generation-url", help="Explicit authorized real local generation endpoint; run one small model workflow only")
    parser.add_argument("--embedding-url", help="Explicit authorized local embedding endpoint")
    parser.add_argument("--model", default="qwen3.6:27b")
    parser.add_argument("--extended-model", action="store_true", help="Add bounded ordinary summaries/report, large QA, mail screen/research/exports")
    parser.add_argument("--scenario", choices=("all", "ordinary", "large-qa"), default="all",
                        help="Run one model scenario in a fresh instance; ordinary/large-qa enable extended mode")
    args = parser.parse_args()
    if args.extended_model and not args.generation_url: parser.error("--extended-model requires --generation-url")
    if args.scenario != "all" and not args.generation_url: parser.error("--scenario requires --generation-url")
    runner = Acceptance(args.evidence, args.timeout, args.generation_url, args.embedding_url, args.model,
                        args.extended_model or args.scenario != "all", args.scenario)
    runner.run()
    return 1 if any(row["status"] == "failed" for row in runner.results) else 0


if __name__ == "__main__":
    raise SystemExit(main())
