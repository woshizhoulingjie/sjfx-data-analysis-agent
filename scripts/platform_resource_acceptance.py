"""Real, synthetic document/OCR/translation/model probes on either host OS.

Run with an explicitly isolated SJFX_ENV_FILE and state/output/temp directories.
No production document, database or credentials are read. Resources are read-only.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import gzip
import importlib.metadata
import io
import json
import math
import os
from pathlib import Path
import platform
import sys
import tarfile
import time
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
TOKEN = "SJFXTEST2026"
SOURCE = TOKEN + " Project budget is 125000 yuan. Delivery date is 2026-12-20."


def samples(folder):
    folder.mkdir(parents=True, exist_ok=True)
    values = {}
    def add(name, content):
        path = folder / name
        path.write_text(content, encoding="utf-8")
        values[name] = path
    for name, content in {
        "中文 文本.TXT": SOURCE + "\n验收材料：预算125000元。",
        "table.csv": "project,budget\n" + TOKEN + ",125000\n",
        "page.html": "<html><body><h1>" + TOKEN + "</h1><p>" + SOURCE + "</p></body></html>",
        "notes.md": "# " + TOKEN + "\n" + SOURCE,
        "data.json": json.dumps({"project": TOKEN, "budget": 125000}),
        "legacy.rtf": r"{\rtf1\ansi " + SOURCE + "}",
        "message.eml": "From: test@example.invalid\nTo: receiver@example.invalid\nDate: Tue, 13 Oct 2026 09:00:00 +0800\nMessage-ID: <sjfx-acceptance@example.invalid>\nSubject: " + TOKEN + "\nMIME-Version: 1.0\nContent-Type: text/plain; charset=utf-8\n\n" + SOURCE,
    }.items(): add(name, content)
    from docx import Document
    doc = Document(); doc.add_heading(TOKEN, 0); doc.add_paragraph(SOURCE)
    doc.add_paragraph("合成验收材料，预算125000元。")
    doc.save(folder / "办公 文档.docx"); values["办公 文档.docx"] = folder / "办公 文档.docx"
    from pptx import Presentation
    slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = TOKEN; slide.placeholders[1].text = SOURCE
    slides.save(folder / "slides.pptx"); values["slides.pptx"] = folder / "slides.pptx"
    from openpyxl import Workbook
    book = Workbook(); book.active.append(["Project", "Budget"]); book.active.append([TOKEN, 125000])
    book.save(folder / "sheet.xlsx"); book.close(); values["sheet.xlsx"] = folder / "sheet.xlsx"
    from PIL import Image, ImageDraw, ImageFont
    fonts = ("C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    font = next((ImageFont.truetype(p, 48) for p in fonts if Path(p).is_file()), ImageFont.load_default(size=48))
    image = Image.new("RGB", (1800, 600), "white"); draw = ImageDraw.Draw(image)
    for i, line in enumerate([TOKEN, "Project budget is 125000 yuan.", "Delivery date is 2026-12-20."]):
        draw.text((70, 80 + i * 115), line, fill="black", font=font)
    image.save(folder / "scan.png"); image.save(folder / "scan.pdf", "PDF", resolution=150)
    values.update({n: folder / n for n in ("scan.png", "scan.pdf")})
    # A small valid text PDF with its own xref table, independent of reportlab.
    text = ("BT /F1 18 Tf 50 700 Td (" + SOURCE + ") Tj ET").encode("ascii")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 900 800] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Length " + str(len(text)).encode() + b" >>\nstream\n" + text + b"\nendstream"]
    raw = b"%PDF-1.4\n"; offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(raw)); raw += str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(raw); raw += b"xref\n0 6\n0000000000 65535 f \n"
    raw += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    raw += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    (folder / "text.pdf").write_bytes(raw); values["text.pdf"] = folder / "text.pdf"
    with zipfile.ZipFile(folder / "archive.zip", "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("中文 文件.txt", SOURCE)
    with tarfile.open(folder / "archive.tar.gz", "w:gz") as t:
        body = SOURCE.encode(); info = tarfile.TarInfo("sample.txt"); info.size = len(body)
        t.addfile(info, io.BytesIO(body))
    with gzip.open(folder / "sample.txt.gz", "wb") as f: f.write(SOURCE.encode())
    values.update({n: folder / n for n in ("archive.zip", "archive.tar.gz", "sample.txt.gz")})
    return values


def main():
    arg = argparse.ArgumentParser()
    arg.add_argument("--work", required=True, type=Path)
    arg.add_argument("--mode", choices=("parse", "translation", "model", "all"), default="all")
    args = arg.parse_args(); args.work.mkdir(parents=True, exist_ok=True)
    if not os.environ.get("SJFX_ENV_FILE"):
        raise SystemExit("Set an explicit isolated SJFX_ENV_FILE before running")
    from config import Config
    report = {"started_at": datetime.now(timezone.utc).isoformat(), "os": platform.platform(),
              "architecture": platform.machine(), "python": platform.python_version(),
              "mode": args.mode, "cases": [], "synthetic_data_only": True}
    report["dependencies"] = {}
    for package in ("torch", "docling", "rapidocr", "onnxruntime", "ctranslate2", "transformers", "pypdf", "rarfile", "py7zr"):
        try: report["dependencies"][package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError: report["dependencies"][package] = None
    def run(name, action):
        row = {"case": name, "verification": "real_execution", "passed": False}; started = time.monotonic()
        try:
            detail = action(); row.update(detail); row["passed"] = bool(detail.get("passed"))
        except Exception as exc:
            row.update(error_type=type(exc).__name__, error=str(exc)[:1600])
        row["elapsed_seconds"] = round(time.monotonic() - started, 3)
        report["cases"].append(row)
        (args.work / "resource-results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(row, ensure_ascii=False), flush=True)
    if args.mode in ("all", "parse"):
        from services.unified_parser import UnifiedDocumentParser
        parser = UnifiedDocumentParser(Config.DOCLING_ARTIFACTS_DIR, Config.RAPIDOCR_MODEL_DIR, max_chars=100000)
        report["parser_status_before"] = parser.status()
        for name, path in samples(args.work / "合成 样本").items():
            def parse(path=path):
                value = parser.parse(path, mode="accurate")
                text = str(value.get("text") or "")
                return {"passed": TOKEN.lower() in text.lower() and value.get("parse_status") == "complete",
                        "content_verified": TOKEN.lower() in text.lower(), "parse_status": value.get("parse_status"),
                        "text": text[:3000], "parser": value.get("parser"), "coverage": value.get("coverage"),
                        "warnings": value.get("warnings")}
            run("parse:" + name, parse)
        report["parser_status_after"] = parser.status()
        def thai_english():
            from services.mail.thai_ocr import ocr_image
            text = ocr_image(args.work / "合成 样本" / "scan.png", time.monotonic() + 40)
            return {"passed": TOKEN.lower() in text.lower(), "text": text, "languages": "tha+eng"}
        run("mail:tesseract_tha_eng", thai_english)
    if args.mode in ("all", "translation"):
        def translate():
            from services.offline_translation import OfflineNLLBProvider
            from services.translation import TranslationService
            provider = OfflineNLLBProvider(Config.TRANSLATION_CT2_MODEL_PATH, cpu_threads=2, max_new_tokens=128)
            value = TranslationService(provider=provider).translate_document({"text": "The project budget is 125000 yuan."})
            text = value.get('translated_text') or ''
            return {"passed": value.get('status') == 'completed' and "125000" in text and any("\u4e00" <= c <= "\u9fff" for c in text),
                    "source": "The project budget is 125000 yuan.", "translation": text,
                    "provider": provider.provider_id, "business_result": value}
        run("translation:nllb_ct2_en_zh", translate)
    if args.mode in ("all", "model"):
        from services.model_backend import create_model_client
        from services.agent_runtime import PydanticAgentRuntime
        client = create_model_client(Config)
        def inference():
            value = client.chat("只依据材料回答问题。", "材料：合成项目预算125000元。问题：预算是多少？", max_tokens=128, timeout=60, retries=0)
            return {"passed": "125000" in value.get("content", ""), "response": value}
        run("model:grounded_generation", inference)
        def structured():
            value = PydanticAgentRuntime(client).chat_json("依据合成材料输出JSON，amount为预算整数。", "预算125000元。", required_fields=["amount"], max_tokens=128, timeout=60, retries=0)
            return {"passed": value.get("json", {}).get("amount") == 125000, "response": value}
        run("model:business_structured_boundary", structured)
        def vector():
            from services.embedding_backend import create_embedding_client
            embed = create_embedding_client(Config)
            vectors = embed.embed(["query: project budget", "passage: The project budget is 125000 yuan.", "passage: Bananas grow in warm climates."])
            def cosine(a,b): return sum(x*y for x,y in zip(a,b)) / math.sqrt(sum(x*x for x in a)*sum(y*y for y in b))
            scores = [cosine(vectors[0], v) for v in vectors[1:]]
            return {"passed": len(vectors) == 3 and all(math.isfinite(v) for row in vectors for v in row) and scores[0] > scores[1],
                    "dimension": len(vectors[0]), "retrieval_scores": scores, "backend": Config.EMBED_BACKEND}
        run("embedding:semantic_retrieval", vector)
    report["passed"] = bool(report["cases"]) and all(row["passed"] for row in report["cases"])
    report["finished_at"] = datetime.now(timezone.utc).isoformat()
    (args.work / "resource-results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
