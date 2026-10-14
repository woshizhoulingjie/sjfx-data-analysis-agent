"""Owner-scoped projections of saved mail result sets."""

from __future__ import annotations

import csv
import html
import json
import math
from pathlib import Path
from uuid import uuid4


def export_result_set(repository, result_set_id, owner_id, output_dir, format_name="csv"):
    result_set = repository.get_result_set(result_set_id, owner_id)
    if not result_set:
        raise ValueError("结果集不存在或不属于当前用户")
    format_name = str(format_name or "csv").casefold()
    if format_name not in {"csv", "html", "xlsx"}:
        raise ValueError("导出格式仅支持 csv、html 或 xlsx")
    messages = []
    for item in result_set["items"]:
        if item["item_type"] != "message":
            continue
        message = repository.get_message(item["item_id"], owner_id, result_set["package_id"])
        if message:
            messages.append({
                "message_key": item["item_id"], "status": item.get("status"),
                "sent_at_utc": message.get("sent_at_utc"), "subject": message.get("subject_raw"),
                "rationale": item.get("rationale") or {},
            })
    fields=(result_set.get("definition") or {}).get("fields") or []
    def safe(value):
        value=str(value if value is not None else "")
        return "'"+value if value.startswith(("=","+","-","@","\t","\r")) else value
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    basename = "mail-research-{}-{}".format(result_set_id[-12:], uuid4().hex[:8])
    if format_name == "csv":
        target = output / (basename + ".csv")
        with target.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["message_key", "status", "sent_at_utc", "subject", "rationale"])
            writer.writeheader()
            for item in messages:
                writer.writerow({k:safe(v) for k,v in item.items()})
            if fields:
                writer.writerow({"subject":"提取字段（含原文与证据）"})
                for field in fields:writer.writerow({"message_key":field.get("message_key"),"status":field.get("status"),"subject":safe(field.get("subject_ref")),"rationale":safe(json.dumps(field,ensure_ascii=False))})
        return target
    if format_name == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise RuntimeError("当前环境未安装 XLSX 导出组件") from exc
        target = output / (basename + ".xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "邮件研究结果"
        sheet.append(["邮件编号", "筛查状态", "发送时间（UTC）", "原文主题", "证据定位"])
        summary = workbook.create_sheet("中文研究结论", 0)
        summary.append(["研究目标", safe((result_set.get("definition") or {}).get("question"))])
        final = result_set.get("final") or {}
        summary.append(["中文结论", safe(final.get("answer") or "此结果集暂无研究结论")])
        for finding in final.get("findings", []):
            summary.append([safe(finding.get("text_zh")), safe(finding.get("evidence_id")), safe(finding.get("source_quote"))])
        for limitation in final.get("limitations", []):
            summary.append(["覆盖缺口", safe(limitation)])
        workbook.active = 0
        for item in messages:
            sheet.append([safe({"matched":"匹配","unknown":"待核查","failed":"失败"}.get(item[key],item[key]) if key=="status" else item[key]) for key in ["message_key","status","sent_at_utc","subject","rationale"]])
        details=workbook.create_sheet("提取字段")
        keys=["subject_ref","fact_type","fact_key","value_raw","value_normalized","status","message_key","evidence_id","quote"]
        details.append(["事项", "字段类型", "字段名", "原文值", "规范值", "状态", "邮件ID", "证据ID", "原文句段", "数值", "币种", "单位", "数量条件"])
        for field in fields:
            value=field.get("value_normalized")
            normalized=dict(value) if isinstance(value,dict) else {}
            number=normalized.get("amount",normalized.get("quantity"))
            numeric=float(number) if number is not None else None
            if numeric is not None and not math.isfinite(numeric):numeric=None
            pretty=(str(number)+" "+str(normalized.get("currency") or normalized.get("unit") or "")).strip() if number is not None else value
            row=[safe(field.get(k)) for k in keys]
            row[0]=safe(field.get("subject_ref") or field.get("subject_raw"))
            row[1]={"amount":"金额","date":"日期","quantity":"数量","owner":"责任方","term":"条件","status":"状态","version":"版本"}.get(field.get("fact_type"),field.get("fact_type"))
            row[4]=safe(pretty)
            details.append(row+[numeric,normalized.get("currency"),normalized.get("unit"),{"at_most":"最多","at_least":"至少","equal":"等于"}.get(normalized.get("comparison"),"")])
            details.cell(details.max_row,10).number_format='#,##0.######' 
        coverage=workbook.create_sheet("覆盖说明");coverage.append(["覆盖数据",safe(json.dumps(result_set.get("coverage"),ensure_ascii=False))])
        for tab in workbook:
            tab.freeze_panes="A2";tab.auto_filter.ref=tab.dimensions
            for column in tab.columns:tab.column_dimensions[column[0].column_letter].width=28
        from openpyxl.styles import Alignment, Font, PatternFill
        summary.column_dimensions["A"].width=42
        summary.column_dimensions["B"].width=85
        summary.column_dimensions["C"].width=85
        for tab in workbook:
            for cell in tab[1]:
                cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="173B4D")
            for row in tab:
                for cell in row:cell.alignment=Alignment(vertical="top",wrap_text=True)
        workbook.save(target)
        return target
    target = output / (basename + ".html")
    rows = "".join(
        "<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
            html.escape(str(item["message_key"])), html.escape(str(item["status"])),
            html.escape(str(item["sent_at_utc"])), html.escape(str(item["subject"])),
        ) for item in messages
    )
    rows += "".join("<tr><td colspan='4'>{}</td></tr>".format(html.escape(json.dumps(f,ensure_ascii=False))) for f in fields)
    final = result_set.get("final") or {}
    escape = lambda value: html.escape(str(value if value is not None else ""))
    answer = final.get("answer") or "此结果集暂无研究结论"
    question = (result_set.get("definition") or {}).get("question") or "未记录研究目标"
    findings = "".join(
        "<li><pre>{}</pre></li>".format(escape(json.dumps(finding, ensure_ascii=False)))
        for finding in final.get("findings", [])
    )
    limitations = "".join("<li>{}</li>".format(escape(item)) for item in final.get("limitations", []))
    coverage = json.dumps({
        "result_set": result_set.get("coverage"),
        "research": final.get("coverage"),
        "complete": final.get("complete"),
    }, ensure_ascii=False, indent=2)
    target.write_text(
        "<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'>"
        "<title>邮件研究结果</title><style>body{{font:16px/1.7 sans-serif;max-width:1100px;margin:32px auto;padding:0 20px}}"
        "pre{{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit}}table{{border-collapse:collapse;width:100%}}"
        "th,td{{border:1px solid #ccd6df;padding:8px;text-align:left;overflow-wrap:anywhere}}</style></head><body>"
        "<h1>邮件研究结果</h1><section><h2>研究目标</h2><pre>{question}</pre></section>"
        "<section><h2>中文研究结论</h2><pre>{answer}</pre></section>"
        "<section><h2>发现与证据</h2><ul>{findings}</ul></section>"
        "<section><h2>限制与资料缺口</h2><ul>{limitations}</ul></section>"
        "<section><h2>邮件记录与提取字段</h2><table><thead><tr><th>邮件</th><th>状态</th><th>时间</th><th>主题</th>"
        "</tr></thead><tbody>{rows}</tbody></table></section>"
        "<section><h2>覆盖说明</h2><pre>{coverage}</pre></section></body></html>".format(
            question=escape(question), answer=escape(answer), findings=findings,
            limitations=limitations, rows=rows, coverage=escape(coverage),
        ), encoding="utf-8",
    )
    return target
