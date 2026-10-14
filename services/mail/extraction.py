"""Bounded first-pass extraction for explicit mail attachments."""

from __future__ import annotations

import io
from pathlib import Path

from services.parse_temp import temporary_parse_directory

from services.file_capabilities import ARCHIVE_EXTENSIONS, DOCUMENT_EXTENSIONS, IMAGE_EXTENSIONS, MAIL_SOURCE_EXTENSIONS, TEXT_EXTENSIONS, file_extension


def _text(payload, charset=None, warnings=None):
    from .parsers.eml import _decode_text
    return _decode_text(payload, charset, warnings if warnings is not None else [])


def _office_preflight(payload):
    """Inspect package headers before an Office library expands XML/images."""
    import zipfile
    from config import Config
    from .policy import MailPolicyError
    # XML object trees use several times their input size. Keep this below the
    # isolated helper's address-space limit even if the general archive budget
    # is much larger for archive inventories.
    budget = min(Config.MAX_ARCHIVE_UNCOMPRESSED_BYTES, Config.MAIL_MAX_ATTACHMENT_BYTES,
                 max(1024 * 1024, Config.MAIL_PARSE_MEMORY_MB * 1024**2 // 8))
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        infos = archive.infolist()
        if len(infos) > Config.MAX_ARCHIVE_ENTRIES:
            raise MailPolicyError("Office 附件包成员数超过解析预算")
        total = 0
        for info in infos:
            if info.is_dir():
                continue
            if info.flag_bits & 1:
                raise MailPolicyError("Office 附件包已加密，无法抽取")
            total += info.file_size
            if total > budget:
                raise MailPolicyError("Office 附件包解压大小超过解析预算")
            if info.file_size and (not info.compress_size or info.file_size / info.compress_size > Config.MAX_ARCHIVE_COMPRESSION_RATIO):
                raise MailPolicyError("Office 附件包压缩比超过解析预算")


def _pdf(payload, max_chars=500000):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(payload))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF 密码保护，无法抽取")
    units = []
    errors = []
    accumulated = 0
    for number, page in enumerate(reader.pages, 1):
        try:
            units.append({"text": page.extract_text() or "", "page_number": number})
        except Exception as exc:
            errors.append({"page_number": number, "error": str(exc)[:200]})
            units.append({"text": "", "page_number": number})
        accumulated += len(units[-1]["text"]) + 2
        if accumulated > max_chars:
            break
    return "\n\n".join(u["text"] for u in units), {"pages": len(reader.pages), "units": units,
        "processed_pages": len(units), "incomplete": len(units) < len(reader.pages),
        "missing_pages": [u["page_number"] for u in units if not u["text"].strip()], "page_errors": errors}


def _docx(payload):
    from docx import Document
    from docx.oxml.ns import qn
    import zipfile
    import xml.etree.ElementTree as ET
    document = Document(io.BytesIO(payload))

    def inline(element):
        if element.tag == qn('w:del'):
            return ''
        if element.tag == qn('w:t'):
            return element.text or ''
        if element.tag == qn('w:tab'):
            return '\t'
        if element.tag in {qn('w:br'), qn('w:cr')}:
            return '\n'
        return ''.join(inline(child) for child in element)

    def blocks(element):
        if element.tag == qn('w:p'):
            value = inline(element)
            return [value] if value else []
        if element.tag == qn('w:tr'):
            # Walk physical XML cells once: row.cells repeats merged cells and
            # cell.text omits nested tables. Recursing preserves both contents.
            cells = []
            def collect(node):
                if node.tag == qn('w:tc'):
                    cells.append('\n'.join(blocks(node)))
                else:
                    for child in node:
                        collect(child)
            collect(element)
            return ['\t'.join(cells)] if cells else []
        if element.tag == qn('w:del'):
            return []
        return [value for child in element for value in blocks(child)]

    units = [{'text': '\n'.join(blocks(document.element.body)), 'document_part': 'body'}]
    incomplete = bool(list(document.element.body.iter(qn('w:altChunk'))))
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        for name in sorted(archive.namelist()):
            if not name.endswith('.xml') or not name.startswith('word/'):
                continue
            component = name[len('word/'):]
            if not (component.startswith(('header', 'footer')) or component in {'footnotes.xml', 'endnotes.xml', 'comments.xml'}):
                continue
            root = ET.fromstring(archive.read(name))
            text = '\n'.join(blocks(root))
            if text:
                units.append({'text': text, 'document_part': component.rsplit('.', 1)[0]})
            incomplete = incomplete or bool(list(root.iter(qn('w:altChunk'))))
    meta = {'units': units, 'incomplete': incomplete}
    if incomplete:
        meta['warnings'] = ['Word 含外部嵌入文本部件，当前文本结果不完整；保留原件']
    return '\n\n'.join(unit['text'] for unit in units), meta


def _xls(payload, max_chars=500000):
    import sys
    vendor=str(Path(__file__).parent/"_vendor")
    if vendor not in sys.path:sys.path.insert(0,vendor)
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("旧版 XLS 需要可选 xlrd 组件") from exc
    workbook=xlrd.open_workbook(file_contents=payload,on_demand=True)
    try:
        from openpyxl.utils import get_column_letter
        units, accumulated = [], 0
        for sheet in workbook.sheets():
            for i in range(sheet.nrows):
                text = "\t".join(str(v) for v in sheet.row_values(i))
                units.append({"text": text, "sheet_name": sheet.name,
                              "cell_range": "A{}:{}{}".format(i+1, get_column_letter(max(1, sheet.ncols)), i+1)})
                accumulated += len(text) + 2
                if accumulated > max_chars:
                    return "\n".join(u["text"] for u in units), {"sheets": workbook.sheet_names(), "units": units, "incomplete": True}
        return "\n".join(u["text"] for u in units), {"sheets": workbook.sheet_names(), "units": units}
    finally:
        workbook.release_resources()


def _spreadsheet(payload, max_chars=500000):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    from openpyxl.utils import get_column_letter
    units = []
    accumulated = 0
    try:
        names = [s.title for s in workbook.worksheets]
        for sheet in workbook.worksheets:
            for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                line = "\t".join(str(cell) if cell is not None else "" for cell in row)
                if line.strip():
                    units.append({"text": line, "sheet_name": sheet.title,
                        "cell_range": "A{}:{}{}".format(number,get_column_letter(max(1,len(row))),number)})
                    accumulated += len(line) + 2
                    if accumulated > max_chars:
                        return "\n".join(u["text"] for u in units), {"sheets": names, "units": units, "incomplete": True}
        return "\n".join(u["text"] for u in units), {"sheets": names, "units": units}
    finally:
        workbook.close()


def _presentation(payload, max_chars):
    from pptx import Presentation
    presentation = Presentation(io.BytesIO(payload))
    units, accumulated = [], 0

    def shape_text(shapes):
        for shape in shapes:
            if getattr(shape, "has_text_frame", False):
                yield shape.text_frame.text
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    yield "\t".join(cell.text for cell in row.cells)
            if hasattr(shape, "shapes"):
                yield from shape_text(shape.shapes)

    incomplete = False
    for number, slide in enumerate(presentation.slides, 1):
        text = "\n".join(shape_text(slide.shapes))
        units.append({"text": text, "page_number": number, "document_part": "slide"})
        accumulated += len(text) + 2
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text if frame is not None else ""
            if notes.strip():
                units.append({"text": notes, "page_number": number, "document_part": "notes"})
                accumulated += len(notes) + 2
        if accumulated > max_chars:
            incomplete = True
            break
    return "\n\n".join(unit["text"] for unit in units), {
        "units": units, "slides": len(presentation.slides), "incomplete": incomplete,
        "warnings": ["已提取幻灯片、表格和备注文字；嵌入图片与对象需另行核验"],
    }


def _legacy(payload, suffix, max_chars):
    from services.legacy_office import extract_legacy
    from config import Config
    with temporary_parse_directory(prefix="sjfx-mail-legacy-") as folder:
        path = Path(folder) / ("attachment" + suffix)
        path.write_bytes(payload)
        value = extract_legacy(path, max_chars=max_chars, seconds=Config.MAIL_PARSE_TIMEOUT_SECONDS)
    text = str(value.get("text") or "")
    metadata = dict(value.get("metadata") or {})
    warnings = list(value.get("warnings") or [])
    if not text.strip():
        return {"status": "failed", "text": "", "parser": value.get("parser") or "legacy-office",
                "warnings": warnings, "coverage_ratio": 0,
                "error": {"type": "LegacyExtractionError", "message": "; ".join(warnings)[:500] or "旧格式附件未获得可读正文"}}
    metadata.update(incomplete=bool(value.get("truncated") or not metadata.get("content_complete", False)), warnings=warnings)
    return _result(text, value.get("parser") or "legacy-office", metadata, max_chars=max_chars)


def _image_metadata(payload):
    from PIL import Image
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(payload)) as image:
            frames = getattr(image, "n_frames", 1)
            image.verify()
    return {"image_frames": frames}


def _result(text, parser, meta=None, *, max_chars=500000):
    meta = meta or {}
    units = meta.get("units") or [{"text": text}]
    bounded = []
    remaining = max_chars
    clipped = False
    for unit in units:
        value = str(unit.get("text") or "")
        if remaining <= 0:
            clipped = True
            break
        if len(value)>remaining:clipped=True
        bounded.append({**unit, "text": value[:remaining],"text_origin":"native","extraction_method":parser})
        remaining -= len(value) + 2
    truncated = clipped or len(text) > max_chars or len(bounded) < len(units)
    missing = meta.get("missing_pages") or []
    warnings = list(meta.get('warnings') or [])
    if truncated: warnings.append("附件文本达到长度上限，保留原文件；文本结果不完整")
    if missing: warnings.append("PDF 以下页缺少可抽取文本，等待 OCR：" + ",".join(map(str,missing)))
    for error in meta.get('page_errors') or []:
        warnings.append('PDF 第 {} 页抽取失败：{}'.format(error['page_number'],error['error']))
    status = "ocr_pending" if missing else "partial" if truncated or meta.get('incomplete') else "ready"
    coverage = (meta.get("pages", 1)-len(missing))/max(1,meta.get("pages", 1))
    if truncated: coverage = min(coverage, max_chars/max(1,len(text)))
    if meta.get('incomplete'): coverage = None
    # Empty scanned pages still keep their page units for OCR, but separators
    # alone are not extracted text and must not inflate extracted_chars.
    combined_text = "\n\n".join(u["text"] for u in bounded)
    if not combined_text.strip():
        combined_text = ""
        if status == "ready":
            status, coverage = "partial", 0
            warnings.append("附件未获得可读正文；保留原件，空结果不能确认文本覆盖")
    return {"status": status, "text": combined_text, "units": bounded,
        "parser": parser, "warnings": warnings, "coverage_ratio": coverage,
        "meta": {k:v for k,v in meta.items() if k != "units"}, "truncated": truncated}


def _extract_attachment(part, *, max_chars=500000):
    """Return text plus a truthful status; OCR is queued separately."""
    payload = part.raw_payload or b""
    content_type = str(part.content_type or "").casefold()
    filename = str(part.filename_raw or "")
    suffix = file_extension(filename)
    # Containers are never treated as text, even with a misleading MIME type.
    if suffix in ARCHIVE_EXTENSIONS or suffix in MAIL_SOURCE_EXTENSIONS or content_type in {
        "application/zip", "application/x-7z-compressed", "application/x-rar-compressed",
        "application/vnd.rar", "application/x-tar", "application/gzip", "message/rfc822",
        "application/vnd.ms-outlook",
    }:
        return {"status": "restricted", "text": "", "parser": "none", "warnings": ["嵌套压缩或邮件附件默认不递归解析；请通过对应导入入口单独处理"]}
    # A recognized filename has a stable route across the workbench. MIME is
    # the fallback for unnamed/unknown parts, not a way to reinterpret an
    # Office package or image as an unrelated parser's input.
    content_type = {
        ".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xlsm": "application/vnd.ms-excel.sheet.macroenabled.12", ".xls": "application/vnd.ms-excel",
        ".doc": "application/msword", ".ppt": "application/vnd.ms-powerpoint", ".rtf": "application/rtf",
    }.get(suffix, "image/attachment" if suffix in IMAGE_EXTENSIONS else content_type)
    # RTF uses text/rtf but its control syntax must be parsed as a document.
    if (content_type.startswith("text/") and suffix not in DOCUMENT_EXTENSIONS - TEXT_EXTENSIONS and content_type != "text/rtf") or suffix in TEXT_EXTENSIONS:
        warnings = []
        byte_limit = max(1, max_chars) * 4 + 4
        text = _text(payload[:byte_limit], getattr(part, "charset", None), warnings)
        if len(payload) > byte_limit:
            warnings.append("附件文本达到读取预算，保留原文件；文本结果不完整")
        if content_type == "text/html" or suffix in {".html", ".htm"}:
            from .html_safety import html_to_translation_text
            text = html_to_translation_text(text)
        result = _result(text, "mail-text", max_chars=max_chars)
        if warnings:
            result["status"] = "partial"
            result["coverage_ratio"] = None
            result["warnings"].extend(warnings)
        return result
    try:
        if content_type == "application/pdf" or suffix == ".pdf":
            text, meta = _pdf(payload, max_chars)
            return _result(text,"pypdf",meta,max_chars=max_chars)
        if "wordprocessingml" in content_type or suffix == ".docx":
            _office_preflight(payload)
            text, meta = _docx(payload)
            return _result(text,"python-docx",meta,max_chars=max_chars)
        if suffix in {".doc", ".ppt", ".rtf"} or content_type in {"application/msword", "application/vnd.ms-powerpoint", "application/rtf", "text/rtf"}:
            legacy_suffix = suffix if suffix in {".doc", ".ppt", ".rtf"} else {
                "application/msword": ".doc", "application/vnd.ms-powerpoint": ".ppt",
                "application/rtf": ".rtf", "text/rtf": ".rtf",
            }[content_type]
            return _legacy(payload, legacy_suffix, max_chars)
        if suffix == ".pptx" or "presentationml" in content_type:
            _office_preflight(payload)
            text, meta = _presentation(payload, max_chars)
            return _result(text, "python-pptx", meta, max_chars=max_chars)
        if suffix == ".xls" or payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1")) and content_type == "application/vnd.ms-excel":
            text,meta=_xls(payload, max_chars)
            return _result(text,"xlrd",meta,max_chars=max_chars)
        if "spreadsheetml" in content_type or content_type in {"application/vnd.ms-excel", "application/vnd.ms-excel.sheet.macroenabled.12"} or suffix in {".xlsx", ".xlsm"}:
            _office_preflight(payload)
            text, meta = _spreadsheet(payload, max_chars)
            return _result(text,"openpyxl",meta,max_chars=max_chars)
        if content_type.startswith("image/") or suffix in IMAGE_EXTENSIONS:
            meta = _image_metadata(payload)
            return {"status": "ocr_pending", "text": "", "parser": "none", "warnings": ["图片附件等待 OCR"], "coverage_ratio": 0, "meta": meta}
        return {"status": "unsupported", "text": "", "parser": "none", "warnings": ["此附件类型尚未支持文本提取"]}
    except ImportError as exc:
        return {"status": "unsupported", "text": "", "parser": "unavailable", "warnings": ["附件解析缺少可选组件：" + str(exc)[:300]], "coverage_ratio": 0}
    except Exception as exc:
        return {"status": "failed", "text": "", "parser": "failed", "warnings": [], "error": {"type": type(exc).__name__, "message": str(exc)[:500]}}


def extract_attachment(part, *, max_chars=500000):
    from config import Config
    if not Config.MAIL_ENABLE_ISOLATED_PARSING:
        return _extract_attachment(part,max_chars=max_chars)
    import subprocess,sys,json,os
    from pathlib import Path
    with temporary_parse_directory(prefix="sjfx-mail-parse-") as folder:
        root=Path(folder);(root/"payload").write_bytes(part.raw_payload or b"")
        (root/"meta.json").write_text(json.dumps({"content_type":part.content_type,"filename_raw":part.filename_raw,"charset":getattr(part,"charset",None),"max_chars":max_chars}),encoding="utf-8")
        env={**os.environ,"OPENBLAS_NUM_THREADS":"1","OMP_NUM_THREADS":"1"}
        try:
            result=subprocess.run([sys.executable,"-m","services.mail.extraction",str(root)],cwd=str(Path(__file__).resolve().parents[2]),capture_output=True,timeout=Config.MAIL_PARSE_TIMEOUT_SECONDS,env=env)
            if result.returncode:raise RuntimeError("附件解析子进程退出 ({})".format(result.returncode))
            return json.loads((root/"result.json").read_text(encoding="utf-8"))
        except Exception as exc:
            return {"status":"failed","text":"","parser":"isolated-mail","warnings":[],"error":{"type":type(exc).__name__,"message":str(exc)[:300]}}


def extract_ocr(part, parser, cancelled=lambda:False):
    from config import Config
    from services.parse_isolation import runner_for
    from .ocr import select_language
    from .text_provenance import OCR_NOTICE
    if parser is None:return {"status":"unsupported","text":"","parser":"unavailable","warnings":["OCR 解析器尚未配置；配置后可手动重试"]}
    suffix=file_extension(part.filename_raw or "")
    if suffix not in IMAGE_EXTENSIONS | {".pdf"}:
        suffix=".pdf" if part.content_type=="application/pdf" else ".png"
    with temporary_parse_directory(prefix="sjfx-mail-ocr-") as folder:
        path=Path(folder)/("attachment"+suffix);path.write_bytes(part.raw_payload or b"")
        try:
            image_meta = _image_metadata(part.raw_payload or b"") if suffix != ".pdf" else {}

            def finish(value):
                frames = image_meta.get("image_frames", 1)
                if frames > 1:
                    value["meta"] = {**(value.get("meta") or {}), **image_meta, "processed_image_frames": 1}
                    value["warnings"] = list(value.get("warnings") or []) + ["多帧图片当前 OCR 仅处理首帧；其余帧未覆盖，保留原件"]
                    if value.get("status") == "ready":
                        value["status"] = "partial"
                    coverage = value.get("coverage_ratio")
                    value["coverage_ratio"] = min(1 / frames, coverage) if coverage is not None else None
                return value

            language=select_language(part.filename_raw or "")
            from .ocr import available
            if available(language):
                import subprocess,sys,json,os
                env={**os.environ,"OPENBLAS_NUM_THREADS":"1","OMP_NUM_THREADS":"1"}
                result=subprocess.run([sys.executable,"-m","services.mail.extraction",str(path),"ocr",language],cwd=str(Path(__file__).resolve().parents[2]),capture_output=True,timeout=Config.MAIL_OCR_TIMEOUT_SECONDS,env=env)
                if result.returncode:raise RuntimeError("隔离 OCR 进程退出 ({})".format(result.returncode))
                return finish(json.loads((Path(folder)/"ocr-result.json").read_text(encoding="utf-8")))
            value=runner_for(parser).parse(path,path.name,mode="accurate",timeout=Config.MAIL_OCR_TIMEOUT_SECONDS,memory_mb=Config.MAIL_PARSE_MEMORY_MB,cancel_check=cancelled)
            text=str(value.get("text") or "")
            return finish({"status":"partial" if text.strip() else "ocr_no_text","text":text,"parser":"isolated-ocr",
                "units":[{"text":text,"text_origin":"ocr","ocr_quality":{"confidence_available":False,"recognition_accuracy":"unverified"}}],
                "warnings":(value.get("warnings") or [])+["此 OCR 解析器未返回逐页定位，不能确认页码与完整覆盖",OCR_NOTICE],"coverage_ratio":None})
        except ImportError as exc:
            return {"status":"unsupported","text":"","parser":"unavailable","warnings":["OCR 缺少可选组件：" + str(exc)[:300]],"coverage_ratio":0}
        except Exception as exc:
            if cancelled():raise
            return {"status":"failed","text":"","parser":"isolated-ocr","error":{"type":type(exc).__name__,"message":str(exc)[:300]}}


if __name__ == "__main__":
    import sys,json
    from types import SimpleNamespace
    from config import Config
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS,(Config.MAIL_PARSE_MEMORY_MB*1024*1024,Config.MAIL_PARSE_MEMORY_MB*1024*1024))
    except (ImportError,ValueError,OSError):
        pass
    if len(sys.argv)>2 and sys.argv[2]=="ocr":
        import time
        from .ocr import extract_result
        path=Path(sys.argv[1]);language=sys.argv[3] if len(sys.argv)>3 else "auto"
        result=extract_result(path,500000,time.monotonic()+Config.MAIL_OCR_TIMEOUT_SECONDS-1,language=language)
        (path.parent/"ocr-result.json").write_text(json.dumps(result,ensure_ascii=False),encoding="utf-8")
        sys.exit(0)
    root=Path(sys.argv[1]);meta=json.loads((root/"meta.json").read_text(encoding="utf-8"))
    value=_extract_attachment(SimpleNamespace(raw_payload=(root/"payload").read_bytes(),content_type=meta["content_type"],filename_raw=meta["filename_raw"],charset=meta.get("charset")),max_chars=meta["max_chars"])
    (root/"result.json").write_text(json.dumps(value,ensure_ascii=False),encoding="utf-8")
