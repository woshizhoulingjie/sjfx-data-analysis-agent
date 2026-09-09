"""Large-package research briefs and exports; no ordinary storage or model calls."""
from __future__ import annotations

import hashlib
import json
import os
import re
import threading
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

SCHEMA = "large-package-report/2.0"
_lock = threading.RLock()


def _text(value, limit=600):
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    if "%PDF-" in value or (value and value.count("\ufffd") / len(value) > .05):
        return ""
    return re.sub(r"\s+", " ", value).strip()[:limit]


def _json(value):
    try:
        data = json.loads(value or "{}")
        return data if isinstance(data, dict) else {}
    except (ValueError, TypeError):
        return {}


def _size(value):
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if value < 1024 or unit == "TB":
            return "{:.1f} {}".format(value, unit)
        value /= 1024


def _atomic(path, data):
    temporary = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
    try:
        temporary.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def report_directory(store, package_id):
    if not re.fullmatch(r"lp-[A-Za-z0-9_-]+", package_id):
        raise ValueError("大数据包编号无效")
    directory = Path(store.db_path).parent / "large_packages" / package_id / "reports"
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _direction(category, data):
    examples = data["examples"]
    if not examples:
        return None
    playbooks = {
        "财务资料": ("财务指标与业务变化核验", ["不同期间的收入、成本和预算有何变化？", "汇总数字能否与明细及对应文件一致？"], ["先核对字段、币种和统计口径，再做同口径比较", "交叉核验汇总表与原始凭证，定位差异来源"]),
        "结构化数据": ("结构化数据质量与指标研究", ["各文件的字段、主键和统计口径是否一致？", "哪些缺失值、异常值或指标变化值得深入核验？"], ["确认完整行数、字段含义及去重规则", "对选定数据做完整质量检查后再汇总或比较"]),
        "合同制度": ("合同条款与制度执行研究", ["关键义务、期限及例外条款分别是什么？", "不同版本或相关文件之间是否存在冲突？"], ["建立条款与来源文件对应表", "核对版本、生效日期和履约证据"]),
        "项目与技术": ("项目进展与技术依据研究", ["项目目标、里程碑和交付结果之间如何对应？", "技术方案的关键假设和风险有什么支撑资料？"], ["按项目及时间整理代表文件", "交叉核验需求、方案、变更记录与交付证据"]),
    }
    title, questions, methods = playbooks.get(category, (
        category + "的主题与证据研究",
        ["这些资料集中讨论哪些主题、对象和问题？", "不同文件对同一事项的描述有何一致或差异？"],
        ["先深度解析代表文件，确认主题与关键术语", "按主题比对原文并建立来源证据表"],
    ))
    keywords = [word for word, _ in data["keywords"].most_common(6)]
    confidence = round(data["confidence"] / max(1, data["files"]), 3)
    return {
        "title": title, "category": category, "priority": "优先核验" if data["models"] else "候选方向",
        "confidence": confidence, "confidence_note": "沿用文件分类平均置信度，仅供线索排序，不代表研究结论可信度。",
        "rationale": "该类别有 {} 个文件，其中 {} 个有可读内容预览、{} 个已有模型摘要。{}建议从所列代表文件开始核验。".format(
            data["files"], data["text_files"], data["models"], "预览关键词包括 {}。".format("、".join(keywords)) if keywords else ""),
        "research_questions": questions, "methods": methods,
        "representative_documents": [item["source_path"] for item in examples],
        "evidence_chain": examples, "basis": "已保存的文件分类与有限内容预览",
        "limitations": ["本项是基于现有内容的研究建议，需结合代表文件全文核验。"],
        "score": data["models"] * 10 + data["text_files"],
    }


def build_report(store, item):
    """One streaming database pass; retain only bounded representative excerpts.

    Counts cover every file, unlike the paginated file browser. Cached snapshots
    are invalidated by file/package revisions and live solely beside the large DB.
    """
    package_id = item["id"]
    directory = report_directory(store, package_id)
    with _lock, store._connect() as conn:
        conn.execute("BEGIN")
        revision = conn.execute("SELECT COUNT(*),MAX(updated_at) FROM files WHERE package_id=?", (package_id,)).fetchone()
        cache_key = hashlib.sha256(json.dumps([SCHEMA, package_id, item.get("updated_at"), list(revision)]).encode()).hexdigest()[:20]
        cached = directory / "overview.json"
        try:
            report = json.loads(cached.read_text(encoding="utf-8"))
            if report.get("revision") == cache_key:
                return report
        except (ValueError, OSError):
            pass
        coverage = Counter()
        groups = {}
        formats = Counter()
        failures, deep_files = [], []
        modified_min, modified_max, total_bytes = None, None, 0
        for row in conn.execute("SELECT * FROM files WHERE package_id=? ORDER BY path", (package_id,)):
            local, deep = _json(row["local_summary"]), _json(row["deep_summary"])
            coverage["inventory_files"] += 1
            total_bytes += int(row["size"] or 0)
            formats[row["extension"] or "无扩展名"] += 1
            modified = int(row["modified_at_ns"] or 0)
            if modified:
                modified_min = modified if modified_min is None else min(modified_min, modified)
                modified_max = modified if modified_max is None else max(modified_max, modified)
            quick_ok = row["quick_status"] == "completed"
            failed = row["quick_status"] == "failed" or row["deep_status"] == "failed" or deep.get("summary_type") == "local_deep_failed"
            deep_ok = row["deep_status"] == "completed" and deep.get("summary_type") != "local_deep_failed"
            model = deep_ok and deep.get("summary_type") == "model"
            text = "" if row["sensitive"] else _text(deep.get("text_preview") or local.get("sample"), 450)
            coverage["quick_completed"] += quick_ok
            coverage["local_summaries"] += bool(local) and quick_ok
            coverage["metadata_summaries"] += bool(local) and quick_ok and not bool(_text(local.get("sample")))
            coverage["deep_completed"] += deep_ok
            coverage["deep_selected"] += row["deep_status"] != "idle"
            coverage["model_summaries"] += model
            coverage["failed_files"] += failed
            coverage["sensitive_files"] += bool(row["sensitive"])
            coverage["quick_pending"] += row["quick_status"] in {"queued", "running", "retry"}
            coverage["unsupported_files"] += quick_ok and not failed and not bool(text) and not bool(row["sensitive"])
            if failed and len(failures) < 50:
                failures.append({"path": row["path"], "error": _text(row["error"] or deep.get("error") or local.get("error"))})
            if deep_ok and len(deep_files) < 20:
                deep_files.append({"path": row["path"], "summary": {
                    "summary_type": deep.get("summary_type"), "summary": _text(deep.get("summary") or text, 1000),
                    "source_sha256": deep.get("source_sha256"), "coverage": deep.get("coverage") or {},
                }})
            category = row["category"] or "未分类"
            group = groups.setdefault(category, {"files": 0, "confidence": 0, "text_files": 0, "models": 0, "examples": [], "keywords": Counter()})
            group["files"] += 1
            group["confidence"] += float(row["confidence"] or 0)
            group["text_files"] += bool(text)
            group["models"] += model
            # Fixed-size sample per category, with completed deep results preferred.
            if text and not failed:
                evidence = {"source_path": row["path"], "source_sha256": row["sha256"], "text": text,
                            "scope": "深度解析正文预览" if deep_ok else "快速解析有限预览", "deep": deep_ok}
                if len(group["examples"]) < 3:
                    group["examples"].append(evidence)
                elif deep_ok:
                    index = next((i for i, value in enumerate(group["examples"]) if not value["deep"]), None)
                    if index is not None:
                        group["examples"][index] = evidence
            for keyword in (local.get("keywords") or [])[:20] if not row["sensitive"] else []:
                word = _text(keyword, 40)
                if word and (word in group["keywords"] or len(group["keywords"]) < 200):
                    group["keywords"][word] += 1
        categories = {key: value["files"] for key, value in sorted(groups.items(), key=lambda pair: (-pair[1]["files"], pair[0]))}
        candidates = sorted(filter(None, (_direction(key, value) for key, value in groups.items())), key=lambda x: (-x["score"], x["title"]))[:5]
        total = coverage["inventory_files"]
        coverage["quick_ratio"] = round(coverage["quick_completed"] / max(1, total), 4)
        coverage["deep_ratio"] = round(coverage["deep_completed"] / max(1, total), 4)
        coverage["model_ratio"] = round(coverage["model_summaries"] / max(1, total), 4)
        coverage["failed_or_unsupported"] = coverage["failed_files"] + coverage["unsupported_files"]
        # Persist explicit zeroes so clients do not infer missing values.
        for key in ("inventory_files", "quick_completed", "deep_completed", "model_summaries", "local_summaries", "metadata_summaries", "deep_selected", "quick_pending", "sensitive_files", "failed_files", "unsupported_files"):
            coverage.setdefault(key, 0)
        limits = [
            "全量统计只描述已盘点文件；快速解析为有限内容预览，不能视为全文分析。",
            "研究方向由已保存的分类和内容预览归纳，不额外调用模型；建议需结合原文核验。",
            "深度解析与模型摘要只覆盖用户选定范围，模型摘要覆盖率为 {:.1%}。".format(coverage["model_ratio"]),
            "本地摘要中有 {} 个仅含元数据，{} 个敏感文件未展示正文。".format(coverage["metadata_summaries"], coverage["sensitive_files"]),
        ]
        if item.get("phase") == "inventory":
            limits.append("盘点仍在进行，当前概览为增量快照，文件总量尚未封存。")
        if coverage["quick_pending"]:
            limits.append("尚有 {} 个文件等待快速解析。".format(coverage["quick_pending"]))
        if coverage["failed_files"]:
            limits.append("有 {} 个文件解析失败；页面和报告最多展示前 50 条失败记录。".format(coverage["failed_files"]))
        if coverage["unsupported_files"]:
            limits.append("有 {} 个非敏感文件尚未提取到可读正文，可能需要 OCR、格式支持或人工复核。".format(coverage["unsupported_files"]))
        findings = [
            "已盘点 {} 个文件，总大小 {}，形成 {} 个智能分类。".format(total, _size(total_bytes), len(categories)),
            "已完成 {} 个本地摘要，其中 {} 个仅含元数据。".format(coverage["local_summaries"], coverage["metadata_summaries"]),
            "选定范围 {} 个文件，已完成深度解析 {} 个、模型摘要 {} 个。".format(coverage["deep_selected"], coverage["deep_completed"], coverage["model_summaries"]),
        ]
        if categories:
            findings.append("主要分类为 " + "、".join("{}（{} 个）".format(key, value) for key, value in list(categories.items())[:5]) + "。")
        basic = ["数据包目录：" + item["root_path"], findings[0], "文件格式：" + "、".join("{} {} 个".format(key, value) for key, value in formats.most_common(10))]
        if modified_min:
            dates = [datetime.fromtimestamp(value / 1e9, timezone.utc).strftime("%Y-%m-%d") for value in (modified_min, modified_max)]
            basic.append("文件修改时间范围：{} 至 {}（不等同于内容事件时间）".format(*dates))
        brief_ready = str(item.get("status") or "") == "completed" and coverage["deep_selected"] > 0 and coverage["deep_completed"] > 0
        report = {
            "schema_version": SCHEMA, "revision": cache_key, "package_id": package_id,
            "generated_at": datetime.now(timezone.utc).isoformat(), "status": item.get("status"), "phase": item.get("phase"),
            "root_path": item["root_path"], "coverage": dict(coverage), "categories": categories,
            "total_bytes": total_bytes, "total_size_human": _size(total_bytes), "formats": dict(formats),
            "deep_files": deep_files, "failures": failures,
            "research_brief": {
                "title": "大数据包情况概览", "available": brief_ready, "ready": brief_ready, "generation_mode": "local_evidence_brief",
                "basic_information": basic, "key_findings": findings,
                "recommended_research_direction": candidates[0] if brief_ready and candidates else {},
                "direction_candidates": candidates[1:] if brief_ready else [], "limitations": limits,
                "empty_direction_reason": "深度摘要完成后生成推荐研究方向。" if not brief_ready else ("暂未提取到足够的可读内容，请先对代表文件进行深度解析。" if not candidates else ""),
            },
            "download_formats": ["docx", "json"],
        }
        _atomic(cached, report)
        return report


def export_report(store, item, file_format="docx"):
    if file_format not in {"docx", "json"}:
        raise ValueError("仅支持 Word 或 JSON 概览下载")
    report = build_report(store, item)
    if not (report.get("research_brief") or {}).get("ready"):
        raise ValueError("深度摘要尚未完成，暂不能下载大数据包研究简报")
    directory = report_directory(store, item["id"])
    target = directory / ("large-package-overview-" + report["revision"] + "." + file_format)
    with _lock:
        if not target.is_file():
            if file_format == "json":
                _atomic(target, report)
            else:
                temporary = target.with_name(target.name + "." + uuid.uuid4().hex + ".tmp")
                try:
                    write_docx(report, temporary)
                    os.replace(temporary, target)
                finally:
                    temporary.unlink(missing_ok=True)
    return target


def write_docx(report, path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.85)
    for name, size in (("Normal", 11), ("Title", 23), ("Heading 1", 16), ("Heading 2", 13), ("List Bullet", 11)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.15
    doc.add_paragraph("大数据包情况概览", "Title")
    brief, coverage = report["research_brief"], report["coverage"]
    doc.add_paragraph("本报告汇总大数据包的盘点规模、智能分类、已保存摘要与后续研究建议，供选择深入研究范围和核验原文使用。")
    doc.add_paragraph("任务编号：" + report["package_id"])
    doc.add_paragraph("生成时间：" + report["generated_at"])
    def items(values):
        for value in values:
            doc.add_paragraph(_text(value, 2000), "List Bullet")
    doc.add_heading("数据包基本信息", 1)
    items(brief["basic_information"])
    doc.add_heading("处理覆盖情况", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("指标", "文件数", "说明")):
        cell.text = label
        shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "DCE6F1"); cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for label, key, note in (("已盘点", "inventory_files", "已发现的物理文件"), ("本地摘要", "local_summaries", "有限预览或元数据"), ("仅元数据", "metadata_summaries", "尚无正文摘要"), ("已选深度范围", "deep_selected", "用户明确选定的文件"), ("深度解析", "deep_completed", "占已盘点文件 {:.1%}".format(coverage["deep_ratio"])), ("模型摘要", "model_summaries", "占已盘点文件 {:.1%}".format(coverage["model_ratio"])), ("解析失败", "failed_files", "需要复核或重试")):
        for cell, value in zip(table.add_row().cells, (label, str(coverage[key]), note)):
            cell.text = value
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge); element.set(qn("w:val"), "single"); element.set(qn("w:sz"), "4"); element.set(qn("w:color"), "D9D9D9"); borders.append(element)
    table._tbl.tblPr.append(borders)
    for row in table.rows:
        for cell, width in zip(row.cells, (1.5, .9, 4.4)):
            cell.width = Inches(width)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                element = OxmlElement("w:" + edge); element.set(qn("w:w"), "90"); element.set(qn("w:type"), "dxa"); margins.append(element)
            cell._tc.get_or_add_tcPr().append(margins)
    doc.add_heading("关键发现", 1)
    items(brief["key_findings"])
    doc.add_heading("推荐研究方向", 1)
    directions = [brief["recommended_research_direction"]] + brief["direction_candidates"]
    if not directions[0]:
        doc.add_paragraph(brief["empty_direction_reason"])
    for index, direction in enumerate(filter(None, directions)):
        doc.add_heading(("首选方向 " if index == 0 else "候选方向 ") + direction["title"], 2)
        doc.add_paragraph(direction["rationale"])
        doc.add_paragraph("研究价值与优先级：" + direction["priority"])
        doc.add_paragraph("建议研究问题")
        items(direction["research_questions"])
        doc.add_paragraph("建议研究方法")
        items(direction["methods"])
        for evidence in direction["evidence_chain"]:
            doc.add_paragraph("代表文件：" + evidence["source_path"])
            doc.add_paragraph(evidence["scope"] + "：" + evidence["text"])
        doc.add_paragraph(direction["confidence_note"])
    doc.add_heading("智能目录分布", 1)
    items(["{}：{} 个文件".format(key, value) for key, value in report["categories"].items()])
    doc.add_heading("选定范围摘要", 1)
    doc.add_paragraph("以下展示最多 20 个已完成深度解析文件的摘要节选；本地摘要与模型摘要分别标注。")
    for item in report["deep_files"]:
        doc.add_heading(item["path"], 2)
        doc.add_paragraph("摘要来源：" + ("模型摘要" if item["summary"]["summary_type"] == "model" else "深度解析本地预览"))
        doc.add_paragraph(item["summary"]["summary"] or "暂无可读摘要")
    if not report["deep_files"]:
        doc.add_paragraph("当前尚无已完成的深度解析摘要。")
    doc.add_heading("限制与待复核事项", 1)
    items(brief["limitations"])
    items([item["path"] + "：" + item["error"] for item in report["failures"]])
    doc.save(str(path))
